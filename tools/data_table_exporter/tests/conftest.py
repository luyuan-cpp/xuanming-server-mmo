"""测试公共夹具:临时目录里现造 xlsx + 现造 ExporterConfig。

刻意不 mock ``openpyxl``:表头解析(5 行元数据 + 列 span)本身就是最容易出错的一环,
用真 xlsx 走一遍 ``read_table`` 才能保证测试和导表器看到的是同一份 schema。
"""

from __future__ import annotations

import sys
from dataclasses import dataclass, field as dc_field
from pathlib import Path

import openpyxl
import pytest

_PKG_ROOT = Path(__file__).resolve().parent.parent
if str(_PKG_ROOT) not in sys.path:
    sys.path.insert(0, str(_PKG_ROOT))

from core.config_loader import ExporterConfig, LangConfig  # noqa: E402
from core.excel_reader import read_table  # noqa: E402
from core.schema import TableSchema  # noqa: E402

# 与 exporter_config.yaml 一致的表头行号
_METADATA_ROWS = {
    "field_name": 1,
    "field_type": 2,
    "owner": 3,
    "options": 4,
    "comment": 5,
}
_DATA_BEGIN_ROW = 6


@dataclass
class Field:
    """一个逻辑字段(可能横跨多列,比如 repeated)。"""
    name: str
    type: str = "uint32"
    owner: str = "server"
    options: str = ""
    comment: str = ""
    span: int = 1


@dataclass
class TableSpec:
    """一张待生成的测试表。"""
    name: str
    fields: list[Field] = dc_field(default_factory=list)
    rows: list[list] = dc_field(default_factory=list)


def write_xlsx(directory: Path, spec: TableSpec) -> Path:
    """按 5 行表头格式把 *spec* 写成一个 .xlsx,返回文件路径。"""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = spec.name

    col = 1
    for f in spec.fields:
        ws.cell(row=_METADATA_ROWS["field_name"], column=col, value=f.name)
        ws.cell(row=_METADATA_ROWS["field_type"], column=col, value=f.type)
        ws.cell(row=_METADATA_ROWS["owner"], column=col, value=f.owner)
        if f.options:
            ws.cell(row=_METADATA_ROWS["options"], column=col, value=f.options)
        if f.comment:
            ws.cell(row=_METADATA_ROWS["comment"], column=col, value=f.comment)
        col += f.span

    total_cols = sum(f.span for f in spec.fields)
    for r, row in enumerate(spec.rows, start=_DATA_BEGIN_ROW):
        for c in range(total_cols):
            value = row[c] if c < len(row) else None
            if value is not None:
                ws.cell(row=r, column=c + 1, value=value)

    directory.mkdir(parents=True, exist_ok=True)
    path = directory / f"{spec.name}.xlsx"
    wb.save(path)
    return path


def make_config(tmp_path: Path) -> ExporterConfig:
    """指向临时目录的最小可用配置(模板目录用仓库里的真模板)。"""
    out = tmp_path / "out"
    return ExporterConfig(
        data_dir=tmp_path / "data",
        data_begin_row=_DATA_BEGIN_ROW,
        metadata_rows=dict(_METADATA_ROWS),
        generated_dir=out,
        json_dir=out / "tables",
        binary_dir=out / "tables",
        proto_dir=out / "proto",
        proto_python_output_dir=out / "proto" / "python",
        state_dir=tmp_path / "state",
        cpp=LangConfig(enabled=True, bit_index_dir=out / "cpp" / "bit_index"),
        go=LangConfig(enabled=True, bit_index_dir=out / "go" / "bit_index"),
        java=LangConfig(enabled=False),
        template_dir=_PKG_ROOT / "templates",
        config_root=tmp_path,
    )


def build_tables(cfg: ExporterConfig, specs: list[TableSpec]) -> list[TableSchema]:
    """把 specs 落成 xlsx 再读回 TableSchema(和导表器走同一条路径)。"""
    schemas: list[TableSchema] = []
    for spec in specs:
        path = write_xlsx(cfg.data_dir, spec)
        schema = read_table(path, cfg)
        assert schema is not None, f"测试表 {spec.name} 解析失败"
        schemas.append(schema)
    return schemas


@pytest.fixture()
def cfg(tmp_path: Path) -> ExporterConfig:
    return make_config(tmp_path)
