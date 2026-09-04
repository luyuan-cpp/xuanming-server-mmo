"""权威 schema 路径的 fail-closed 回归。

这些用例守的是同一条纪律:**schema 与源表对不上时,导表必须停,不许静默产出。**
每一条都对应一次真实被证实过的静默失败形态,别因为「看起来不会有人这么写」就删。

刻意在 tmp_path 里现造 xlsx + 现造 schema proto,不依赖仓库里的 data/*.xlsx ——
那样测试会随策划改表而红(见 .github/workflows/exporter-tests.yml 的约定)。
"""

from __future__ import annotations

import sys
from pathlib import Path

import openpyxl
import pytest

_PKG_ROOT = Path(__file__).resolve().parent.parent
if str(_PKG_ROOT) not in sys.path:
    sys.path.insert(0, str(_PKG_ROOT))

from core.schema_proto import (  # noqa: E402
    SchemaProtoError,
    build_table_schema,
    load_vocabulary,
)

# 词表是运行时从 cfg_options.proto 解析的,测试也得有一份。
_VOCAB = (_PKG_ROOT.parent.parent / "data" / "schema" / "cfg_options.proto")

_BASE_SCHEMA = '''\
syntax = "proto3";
package mmorpg.cfgtable.v1;
import "cfg_options.proto";

message DemoSub {
  uint32 pack_type = 1;
  uint32 pack_value = 2;
}

message DemoTable {
  option (cfg_sheet)       = "Demo";
  option (cfg_primary_key) = "id";

  uint32 id = 1;
  // 等级
  uint32 level = 2 [(cfg_index) = true];
  repeated uint32 effect = 3 [(cfg_slots) = 3];
  map<string, string> tag = 4 [(cfg_slots) = 2];
  map<string, bool> flags = 5 [(cfg_slots) = 2, (cfg_shape) = CFG_SHAPE_SET];
  repeated DemoSub pack = 6 [(cfg_slots) = 2];
  int32 designer_note = 9001 [(cfg_owner) = "designer"];
}
'''

# 列序必须与 schema 里每个字段的 (cfg_slots) 对得上。
_HEADERS = (["id", "level"]
            + ["effect", "", ""]
            + ["tag", "", "", ""]
            + ["flags", ""]
            + ["pack", "", "", ""]
            + ["designer_note"])


@pytest.fixture
def env(tmp_path: Path):
    """造出 (schema 目录, xlsx 路径),两边一开始是自洽的。"""
    schema_dir = tmp_path / "schema"
    schema_dir.mkdir()
    (schema_dir / "cfg_options.proto").write_bytes(_VOCAB.read_bytes())
    (schema_dir / "demo_table.proto").write_text(_BASE_SCHEMA, encoding="utf-8")

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Demo"
    for i, name in enumerate(_HEADERS, start=1):
        if name:
            ws.cell(row=1, column=i, value=name)
    for r in (6, 7):
        for c in range(1, len(_HEADERS) + 1):
            ws.cell(row=r, column=c, value=1)
    xlsx = tmp_path / "Demo.xlsx"
    wb.save(xlsx)
    return schema_dir, xlsx


def _mutate(schema_dir: Path, old: str, new: str) -> None:
    p = schema_dir / "demo_table.proto"
    text = p.read_text(encoding="utf-8")
    assert old in text, "用例写错了:%r 不在 schema 里" % old
    p.write_text(text.replace(old, new, 1), encoding="utf-8")


def test_baseline_builds(env):
    schema_dir, xlsx = env
    schema = build_table_schema(schema_dir / "demo_table.proto", xlsx)
    assert schema.name == "Demo"
    assert len(schema.columns) == len(_HEADERS)
    # 被 owner 丢弃的列仍要占物理列位,否则 read_data_rows 会按下标错位。
    assert [c.name for c in schema.columns][-1] == "designer_note"
    assert schema.field_numbers["level"] == 2
    assert "designer_note" not in schema.field_numbers


def test_vocabulary_is_parsed_from_proto(env):
    """option 词表来自 cfg_options.proto,不是 Python 侧硬编码。"""
    vocab = load_vocabulary(env[0] / "cfg_options.proto")
    assert "cfg_slots" in vocab.field_options
    assert vocab.field_options["cfg_expr_param"].repeated is True
    assert "CFG_SHAPE_SET" in vocab.enums["CfgShape"]


@pytest.mark.parametrize("old,new,needle", [
    # option 名拼错 -> 从前静默忽略,该列悄悄按默认 owner 进产物
    ('(cfg_index) = true', '(cfg_indx) = true', "未知 option"),
    # 带引号的 "false" 是真值 -> 从前会把这一列当成 key
    ('(cfg_index) = true', '(cfg_index) = "false"', "必须是 true/false"),
    # 容量与实际列数对不上
    ('repeated uint32 effect = 3 [(cfg_slots) = 3]',
     'repeated uint32 effect = 3 [(cfg_slots) = 2]', "(cfg_slots)"),
    # 非标量漏写容量
    ('repeated uint32 effect = 3 [(cfg_slots) = 3]', 'repeated uint32 effect = 3',
     "必须写 (cfg_slots)"),
    # map<K,bool> 不表态到底是 MAP 还是 SET
    ('map<string, bool> flags = 5 [(cfg_slots) = 2, (cfg_shape) = CFG_SHAPE_SET]',
     'map<string, bool> flags = 5 [(cfg_slots) = 2]', "必须显式写 (cfg_shape)"),
    # repeated 一个既不是标量也不是子消息的类型 -> 从前静默降级成标量数组
    ('repeated DemoSub pack', 'repeated Nope pack', "既不是内置标量"),
    # 字段名重复 -> 从前后者静默顶掉前者
    ('uint32 level = 2', 'uint32 id = 2', "字段名 id 重复"),
    # 字段号重复
    ('uint32 level = 2', 'uint32 level = 1', "字段号 1 重复"),
    # protobuf 保留段
    ('uint32 level = 2', 'uint32 level = 19001', "保留段"),
    # 主键不在第一列
    ('option (cfg_primary_key) = "id"', 'option (cfg_primary_key) = "level"',
     "(cfg_primary_key)"),
    # 子消息字段号必须是 1..N —— 号决定它对应哪个子列,不是行序
    ('uint32 pack_type = 1', 'uint32 pack_type = 3', "连续序列"),
    # 块注释会吞掉后面的字段声明
    ('  uint32 id = 1;', '  /* 吞 */\n  uint32 id = 1;', "不支持块注释"),
    # 认不出的行一律报错,不许静默跳过
    ('  uint32 id = 1;', '  uint32 id = 1;\n  乱码一行', "无法解析的行"),
    # 删字段后 reserved 防复用
    ('uint32 level = 2', 'reserved 2;\n  uint32 level = 2', "已 reserved 的号"),
])
def test_fail_closed(env, old, new, needle):
    schema_dir, xlsx = env
    _mutate(schema_dir, old, new)
    with pytest.raises(SchemaProtoError) as exc:
        build_table_schema(schema_dir / "demo_table.proto", xlsx)
    assert needle in str(exc.value)


def test_header_mismatch_is_fail_closed(env):
    """列名对不上 -> 整批不产出。绑定按列名做,所以改名就是改绑定。"""
    schema_dir, xlsx = env
    _mutate(schema_dir, "uint32 level = 2", "uint32 lvl = 2")
    with pytest.raises(SchemaProtoError) as exc:
        build_table_schema(schema_dir / "demo_table.proto", xlsx)
    msg = str(exc.value)
    assert "期望表头 'lvl'" in msg          # schema 侧多出来的
    assert "没有对应字段" in msg             # xlsx 侧没人认领的


def test_field_order_in_proto_is_irrelevant(env):
    """按列名绑定的直接后果:主消息字段声明顺序随便换,结果一模一样。"""
    from core.schema_parity import diff
    schema_dir, xlsx = env
    before = build_table_schema(schema_dir / "demo_table.proto", xlsx)
    _mutate(schema_dir,
            "  uint32 id = 1;\n  // 等级\n  uint32 level = 2 [(cfg_index) = true];",
            "  // 等级\n  uint32 level = 2 [(cfg_index) = true];\n  uint32 id = 1;")
    after = build_table_schema(schema_dir / "demo_table.proto", xlsx)
    assert diff(before, after) == []


def test_positional_check_is_opt_in(env):
    """字段号是权威的:日常导表不再拿历史发号规则去卡它,但迁移期可以打开核对。"""
    schema_dir, xlsx = env
    _mutate(schema_dir, "uint32 level = 2", "uint32 level = 8")
    schema = build_table_schema(schema_dir / "demo_table.proto", xlsx)
    assert schema.field_numbers["level"] == 8      # 默认放行,产物随之改变
    with pytest.raises(SchemaProtoError) as exc:
        build_table_schema(schema_dir / "demo_table.proto", xlsx, check_positional=True)
    assert "按历史发号规则应为 2" in str(exc.value)


def test_excluded_column_number_range(env):
    """不进产物的列必须用 >= 9000 的号,反之亦然。"""
    schema_dir, xlsx = env
    _mutate(schema_dir, "int32 designer_note = 9001", "int32 designer_note = 7")
    with pytest.raises(SchemaProtoError) as exc:
        build_table_schema(schema_dir / "demo_table.proto", xlsx)
    assert "不进产物" in str(exc.value)


def test_comment_roundtrip_preserves_indentation(env):
    """注释是 xlsx 第 5 行原文的容器,缩进必须逐字符还原。

    Buff.buff_type 的注释是一整段带缩进的 C++ enum;从前解析器 .strip() 掉缩进,
    往返就不再等价。
    """
    schema_dir, xlsx = env
    _mutate(schema_dir, "  // 等级\n",
            "  // enum E {\n  //     kA = 0,   // 说明\n  // };\n")
    schema = build_table_schema(schema_dir / "demo_table.proto", xlsx)
    level = next(c for c in schema.columns if c.name == "level")
    assert level.comment == "enum E {\n    kA = 0,   // 说明\n};"


def test_multi_primary_key_forbids_bit_index(env):
    """主键可重复 + 位序 = 无定义,必须互斥。

    位序是「ID -> 存档位图下标」,重复 ID 没有唯一下标。
    """
    schema_dir, xlsx = env
    _mutate(schema_dir, "uint32 id = 1;",
            "uint32 id = 1 [(cfg_multi) = true, (cfg_bit_index) = true];")
    with pytest.raises(SchemaProtoError) as exc:
        build_table_schema(schema_dir / "demo_table.proto", xlsx)
    assert "(cfg_bit_index)" in str(exc.value)


def test_multi_primary_key_flag(env):
    """主键上的 (cfg_multi) 是「id 可重复」的唯一来源,不设第二个表级开关。"""
    schema_dir, xlsx = env
    assert build_table_schema(schema_dir / "demo_table.proto", xlsx).multi_primary_key is False
    _mutate(schema_dir, "uint32 id = 1;", "uint32 id = 1 [(cfg_multi) = true];")
    assert build_table_schema(schema_dir / "demo_table.proto", xlsx).multi_primary_key is True
