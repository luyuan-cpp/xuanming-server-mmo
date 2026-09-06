"""tip 码轴回归测试。

守的红线（说明见 core/generators/enum_gen.py 的模块 docstring）:

1. **段是发号器的输入** —— 新码只能落在本组段内，不是全局队尾。
   这是本次改造的根因修复:改造前往 common 组加一个码会拿到队尾的号，
   落在 common 段外，而全程零报错。
2. **号只增不减** —— 从 xlsx 删掉一行不回收它的号，避免号被复用后
   老客户端把新含义当旧含义。
3. **段满即中止** —— 不静默溢出到下一段。
4. **枚举名全局唯一** —— 生成的 tip proto 无 package 声明，枚举值同处
   一个命名空间，重名会在 protoc 阶段炸，这里提前拦。
5. **v1 state 一律拒绝** —— 扁平队尾发号的旧状态不许被沿用。
6. **段声明缺失一律拒绝** —— 不许有「没有段」的组。
7. **故障分类在表里** —— fault 列按表头名定位、必须存在、取值闭集；
   产物 faults.go 是 serverbase 判「故障 / 业务拒绝」的唯一依据。
"""

from __future__ import annotations

import json
from pathlib import Path

import openpyxl
import pytest

from conftest import Field, TableSpec, make_config, write_xlsx
from core.excel_reader import read_table
from core.generators.enum_gen import (
    TIP_STATE_SCHEMA_VERSION,
    TipAxisError,
    _check_tip_state_evolution,
    generate_tip_enums,
    validate_tip_references,
)

# Tip.xlsx 的数据区从第 18 行开始（与 exporter 一致）
_TIP_DATA_BEGIN = 18


def write_tip_xlsx(
    path: Path,
    groups: list[tuple[str, list[tuple]]],
    *,
    fault_header: str | None = "fault",
    fault_column: int = 3,
) -> Path:
    """按真实 Tip.xlsx 的形状落一个测试用表。

    *groups* 形如 ``[("//g base=1000 width=100", [("CodeName", "文案"), ...]), ...]``；
    组头行原样写入 A 列，方便直接构造「缺 base」这类坏输入。
    entry 也可以是三元组 ``(name, text, fault)``，第三项原样写进 fault 列
    （bool / int / str 都行，用来测取值闭集）。

    *fault_header* 传 None 就不写 fault 表头（测「列缺失必须中止」）；
    *fault_column* 指定 fault 列的 1-based 位置（测「按表头名不按列位」）。
    """
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.cell(row=1, column=1, value="name")
    ws.cell(row=1, column=2, value="commen")
    ws.cell(row=2, column=1, value="string")
    ws.cell(row=2, column=2, value="string")
    ws.cell(row=4, column=1, value="common")
    ws.cell(row=4, column=2, value="design")
    if fault_header is not None:
        ws.cell(row=1, column=fault_column, value=fault_header)
        ws.cell(row=2, column=fault_column, value="bool")
        ws.cell(row=4, column=fault_column, value="server")

    row = _TIP_DATA_BEGIN
    for header, entries in groups:
        ws.cell(row=row, column=1, value=header)
        row += 1
        for entry in entries:
            name, text = entry[0], entry[1]
            ws.cell(row=row, column=1, value=name)
            if text:
                ws.cell(row=row, column=2, value=text)
            if len(entry) > 2 and entry[2] is not None:
                ws.cell(row=row, column=fault_column, value=entry[2])
            row += 1

    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path


def _faults_go(cfg) -> str:
    return (Path(cfg.go.code_dir).parent / "tip" / "faults.go").read_text(encoding="utf-8")


def _state_path(cfg) -> Path:
    return cfg.state_dir / "mapping" / "tip_enum_ids" / "tip_enum_ids.json"


def _read_state(cfg) -> dict:
    return json.loads(_state_path(cfg).read_text(encoding="utf-8"))


def _make_cfg(tmp_path: Path):
    cfg = make_config(tmp_path)
    cfg.tip_file = tmp_path / "data" / "tip" / "Tip.xlsx"
    cfg.go.code_dir = tmp_path / "out" / "go" / "generated" / "table"
    cfg.tip_axis_bootstrap = True
    return cfg


def test_missing_tip_file_is_rejected(tmp_path: Path):
    cfg = _make_cfg(tmp_path)

    with pytest.raises(TipAxisError) as exc:
        generate_tip_enums(cfg)
    assert "唯一权威源" in str(exc.value) and str(cfg.tip_file) in str(exc.value)


def test_missing_state_requires_explicit_bootstrap(tmp_path: Path):
    cfg = _make_cfg(tmp_path)
    cfg.tip_axis_bootstrap = False
    write_tip_xlsx(cfg.tip_file, [
        ("//alpha_error base=1000 width=100", [("A1", "")]),
    ])

    with pytest.raises(TipAxisError) as exc:
        generate_tip_enums(cfg)
    assert "--tip-axis-bootstrap" in str(exc.value)
    assert not cfg.proto_dir.exists(), "state 缺失时不能留下任何生成产物"


def test_v2_state_evolution_rejects_deleted_tombstone():
    previous = {
        "schema_version": 2,
        "groups": {
            "alpha_error": {
                "base": 1000,
                "width": 100,
                "ids": {"A1": 1000, "Retired": 1001},
            }
        },
        "legacy_ids": {},
    }
    current = {
        "schema_version": 2,
        "groups": {
            "alpha_error": {
                "base": 1000,
                "width": 100,
                "ids": {"A1": 1000},
            }
        },
        "legacy_ids": {},
    }

    with pytest.raises(TipAxisError) as exc:
        _check_tip_state_evolution(previous, current)
    assert "Retired" in str(exc.value) and "只增不减" in str(exc.value)


def test_v2_state_evolution_allows_new_ids_and_segment_growth():
    previous = {
        "schema_version": 2,
        "groups": {
            "alpha_error": {"base": 1000, "width": 100, "ids": {"A1": 1000}}
        },
        "legacy_ids": {"1": {"group": "alpha_error", "name": "A1", "new": 1000}},
    }
    current = {
        "schema_version": 2,
        "groups": {
            "alpha_error": {
                "base": 900,
                "width": 300,
                "ids": {"A1": 1000, "A2": 1001},
            }
        },
        "legacy_ids": {"1": {"group": "alpha_error", "name": "A1", "new": 1000}},
    }

    _check_tip_state_evolution(previous, current)


def test_v1_to_v2_requires_complete_legacy_mapping():
    previous = {"alpha_error": {"A1": 1}}
    current = {
        "schema_version": 2,
        "groups": {
            "alpha_error": {"base": 1000, "width": 100, "ids": {"A1": 1000}}
        },
        "legacy_ids": {},
    }

    with pytest.raises(TipAxisError) as exc:
        _check_tip_state_evolution(previous, current)
    assert "迁移记录不完整" in str(exc.value)

    current["legacy_ids"] = {
        "1": {"group": "alpha_error", "name": "A1", "new": 1000}
    }
    _check_tip_state_evolution(previous, current)


# ---------------------------------------------------------------- 红线 1

def test_new_code_lands_in_its_own_segment(tmp_path: Path):
    """改造前的真实故障:往第一组加码会拿到全局队尾的号。

    两组，第一组段 [1000,1100)，第二组段 [2000,2100)。
    先各发一批，再往**第一组**追加一个新码 —— 它必须拿 1000 段里的空位，
    而不是接在第二组已用最大号后面。
    """
    cfg = _make_cfg(tmp_path)
    write_tip_xlsx(cfg.tip_file, [
        ("//alpha_error base=1000 width=100", [("AlphaOne", "一"), ("AlphaTwo", "二")]),
        ("//beta_error base=2000 width=100", [("BetaOne", "一"), ("BetaTwo", "二")]),
    ])
    generate_tip_enums(cfg)
    first = _read_state(cfg)["groups"]
    assert first["alpha_error"]["ids"] == {"AlphaOne": 1000, "AlphaTwo": 1001}
    assert first["beta_error"]["ids"] == {"BetaOne": 2000, "BetaTwo": 2001}

    # 往 alpha 追加一个码
    write_tip_xlsx(cfg.tip_file, [
        ("//alpha_error base=1000 width=100",
         [("AlphaOne", "一"), ("AlphaTwo", "二"), ("AlphaThree", "三")]),
        ("//beta_error base=2000 width=100", [("BetaOne", "一"), ("BetaTwo", "二")]),
    ])
    generate_tip_enums(cfg)
    ids = _read_state(cfg)["groups"]["alpha_error"]["ids"]
    assert ids["AlphaThree"] == 1002, "新码必须落在本组段内,而不是全局队尾"
    # 老号一个都不许动
    assert ids["AlphaOne"] == 1000 and ids["AlphaTwo"] == 1001


# ---------------------------------------------------------------- 红线 2

def test_removed_code_keeps_its_id_and_is_not_reused(tmp_path: Path):
    """删掉一行不回收号:新码不能捡走被删码的号。"""
    cfg = _make_cfg(tmp_path)
    write_tip_xlsx(cfg.tip_file, [
        ("//alpha_error base=1000 width=100",
         [("AlphaOne", "一"), ("AlphaTwo", "二"), ("AlphaThree", "三")]),
    ])
    generate_tip_enums(cfg)
    assert _read_state(cfg)["groups"]["alpha_error"]["ids"]["AlphaTwo"] == 1001

    # 删掉中间那个，再加一个新的
    write_tip_xlsx(cfg.tip_file, [
        ("//alpha_error base=1000 width=100",
         [("AlphaOne", "一"), ("AlphaThree", "三"), ("AlphaFour", "四")]),
    ])
    generate_tip_enums(cfg)
    ids = _read_state(cfg)["groups"]["alpha_error"]["ids"]
    assert ids["AlphaTwo"] == 1001, "被删的码要留在 state 里当墓碑"
    assert ids["AlphaFour"] != 1001, "新码不许捡走被删码的号"
    assert ids["AlphaFour"] == 1003


# ---------------------------------------------------------------- 红线 3

def test_segment_full_aborts(tmp_path: Path):
    """段满不静默溢出到下一段。"""
    cfg = _make_cfg(tmp_path)
    write_tip_xlsx(cfg.tip_file, [
        ("//alpha_error base=1000 width=2", [("A1", ""), ("A2", ""), ("A3", "")]),
    ])
    with pytest.raises(TipAxisError) as exc:
        generate_tip_enums(cfg)
    assert "已满" in str(exc.value) and "A3" in str(exc.value)


def test_overlapping_segments_abort(tmp_path: Path):
    cfg = _make_cfg(tmp_path)
    write_tip_xlsx(cfg.tip_file, [
        ("//alpha_error base=1000 width=1000", [("A1", "")]),
        ("//beta_error base=1500 width=1000", [("B1", "")]),
    ])
    with pytest.raises(TipAxisError) as exc:
        generate_tip_enums(cfg)
    assert "段重叠" in str(exc.value)


# ---------------------------------------------------------------- 红线 4

def test_duplicate_enum_name_across_groups_aborts(tmp_path: Path):
    """tip proto 无 package,枚举值全局同名空间,重名要在 protoc 之前拦住。"""
    cfg = _make_cfg(tmp_path)
    write_tip_xlsx(cfg.tip_file, [
        ("//alpha_error base=1000 width=100", [("Duplicated", "一")]),
        ("//beta_error base=2000 width=100", [("Duplicated", "二")]),
    ])
    with pytest.raises(TipAxisError) as exc:
        generate_tip_enums(cfg)
    assert "Duplicated" in str(exc.value)


# ---------------------------------------------------------------- 红线 5

def test_v1_state_is_rejected(tmp_path: Path):
    """v1 是扁平队尾发号的产物,沿用它等于把根因带进来。"""
    cfg = _make_cfg(tmp_path)
    write_tip_xlsx(cfg.tip_file, [
        ("//alpha_error base=1000 width=100", [("A1", "")]),
    ])
    p = _state_path(cfg)
    p.parent.mkdir(parents=True, exist_ok=True)
    p.write_text(json.dumps({"alpha_error": {"A1": 1}}), encoding="utf-8")

    with pytest.raises(TipAxisError) as exc:
        generate_tip_enums(cfg)
    assert "v1" in str(exc.value)


def test_corrupted_state_is_rejected_instead_of_reset(tmp_path: Path):
    """已有 state 损坏时必须 fail-closed，不能当成首次运行重新发号。"""
    cfg = _make_cfg(tmp_path)
    write_tip_xlsx(cfg.tip_file, [
        ("//alpha_error base=1000 width=100", [("A1", "")]),
    ])
    p = _state_path(cfg)
    p.parent.mkdir(parents=True, exist_ok=True)
    p.write_text("{not valid json", encoding="utf-8")

    with pytest.raises(TipAxisError) as exc:
        generate_tip_enums(cfg)
    assert "读取失败" in str(exc.value) and str(p) in str(exc.value)
    assert p.read_text(encoding="utf-8") == "{not valid json"
    assert not cfg.proto_dir.exists(), "state 失败后不能留下任何生成产物"


def test_state_io_error_is_rejected(tmp_path: Path):
    """state 路径不可读时不能被误判成文件不存在。"""
    cfg = _make_cfg(tmp_path)
    write_tip_xlsx(cfg.tip_file, [
        ("//alpha_error base=1000 width=100", [("A1", "")]),
    ])
    p = _state_path(cfg)
    p.mkdir(parents=True)

    with pytest.raises(TipAxisError) as exc:
        generate_tip_enums(cfg)
    assert "读取失败" in str(exc.value) and str(p) in str(exc.value)


def test_existing_empty_state_is_not_treated_as_first_run(tmp_path: Path):
    cfg = _make_cfg(tmp_path)
    write_tip_xlsx(cfg.tip_file, [
        ("//alpha_error base=1000 width=100", [("A1", "")]),
    ])
    p = _state_path(cfg)
    p.parent.mkdir(parents=True, exist_ok=True)
    p.write_text("{}", encoding="utf-8")

    with pytest.raises(TipAxisError) as exc:
        generate_tip_enums(cfg)
    assert "v1" in str(exc.value)


def test_removing_or_renaming_whole_group_is_rejected(tmp_path: Path):
    """整组消失会丢掉段墓碑，必须显式保留空组头而不是静默复用。"""
    cfg = _make_cfg(tmp_path)
    write_tip_xlsx(cfg.tip_file, [
        ("//alpha_error base=1000 width=100", [("A1", "")]),
    ])
    generate_tip_enums(cfg)

    # 即使新组使用完全不重叠的段也要拒绝，证明守的是整组连续性，
    # 不是碰巧被「段重叠」挡住。
    write_tip_xlsx(cfg.tip_file, [
        ("//beta_error base=2000 width=100", [("B1", "")]),
    ])
    with pytest.raises(TipAxisError) as exc:
        generate_tip_enums(cfg)
    assert "alpha_error" in str(exc.value) and "空组头" in str(exc.value)


def test_retired_group_keeps_empty_header_and_allocated_tombstones(tmp_path: Path):
    cfg = _make_cfg(tmp_path)
    write_tip_xlsx(cfg.tip_file, [
        ("//alpha_error base=1000 width=100", [("A1", "")]),
    ])
    generate_tip_enums(cfg)

    write_tip_xlsx(cfg.tip_file, [
        ("//alpha_error base=1000 width=100", []),
        ("//beta_error base=2000 width=100", [("B1", "")]),
    ])
    generate_tip_enums(cfg)

    state = _read_state(cfg)
    assert state["groups"]["alpha_error"]["ids"] == {"A1": 1000}
    segments = (Path(cfg.go.code_dir).parent / "tip" / "segments.go").read_text(encoding="utf-8")
    assert 'Domain: "alpha"' in segments and "Lo: 1000, Hi: 1000, Count: 1" in segments


def test_id_outside_declared_segment_aborts(tmp_path: Path):
    """有人改了 base 却没迁老号 —— 必须报出来,不能默默按新段发号。"""
    cfg = _make_cfg(tmp_path)
    write_tip_xlsx(cfg.tip_file, [
        ("//alpha_error base=5000 width=100", [("A1", "")]),
    ])
    p = _state_path(cfg)
    p.parent.mkdir(parents=True, exist_ok=True)
    p.write_text(json.dumps({
        "schema_version": TIP_STATE_SCHEMA_VERSION,
        "groups": {"alpha_error": {"base": 1000, "width": 100, "ids": {"A1": 1000}}},
        "legacy_ids": {},
    }, ensure_ascii=False), encoding="utf-8")

    with pytest.raises(TipAxisError) as exc:
        generate_tip_enums(cfg)
    assert "段外" in str(exc.value)


# ---------------------------------------------------------------- 红线 6

def test_group_header_without_segment_aborts(tmp_path: Path):
    """旧格式的组头(只有组名)一律拒绝 —— 没有段就没有分配规则。"""
    cfg = _make_cfg(tmp_path)
    write_tip_xlsx(cfg.tip_file, [
        ("//alpha_error", [("A1", "")]),
    ])
    with pytest.raises(TipAxisError) as exc:
        generate_tip_enums(cfg)
    assert "段声明" in str(exc.value)


def test_plain_comment_line_does_not_break_grouping(tmp_path: Path):
    """组内的普通注释行不该把后续的码甩出组外。"""
    cfg = _make_cfg(tmp_path)
    write_tip_xlsx(cfg.tip_file, [
        ("//alpha_error base=1000 width=100", [("A1", "")]),
    ])
    # 手工插一行不带 base 的注释，再补一个码
    wb = openpyxl.load_workbook(cfg.tip_file)
    ws = wb.active
    ws.cell(row=_TIP_DATA_BEGIN + 2, column=1, value="// 这是组内说明,不是组头")
    ws.cell(row=_TIP_DATA_BEGIN + 3, column=1, value="A2")
    wb.save(cfg.tip_file)

    generate_tip_enums(cfg)
    ids = _read_state(cfg)["groups"]["alpha_error"]["ids"]
    assert ids == {"A1": 1000, "A2": 1001}


# ---------------------------------------------------------------- 产物

def test_text_table_and_segment_table_are_generated(tmp_path: Path):
    """B 列(中文文案)必须有出口 —— 改造前它从没被读过。"""
    cfg = _make_cfg(tmp_path)
    write_tip_xlsx(cfg.tip_file, [
        ("//alpha_error base=1000 width=100", [("A1", "甲一"), ("A2", "")]),
    ])
    generate_tip_enums(cfg)

    text = json.loads((Path(cfg.json_dir) / "tip_text.json").read_text(encoding="utf-8"))
    assert text == {"1000": "甲一"}, "有文案的码进表,没文案的不编造"

    seg = (Path(cfg.go.code_dir).parent / "tip" / "segments.go").read_text(encoding="utf-8")
    assert "package tip" in seg
    assert 'Domain: "alpha"' in seg and "Base: 1000" in seg


def test_tip_proto_declares_java_package(tmp_path: Path):
    """Tip Java 绑定必须落入实际部署的 com/game/table 包，而非默认包孤岛。"""
    cfg = _make_cfg(tmp_path)
    cfg.java.enabled = True
    cfg.java.package = "com.game.table"
    write_tip_xlsx(cfg.tip_file, [
        ("//alpha_error base=1000 width=100", [("A1", "甲一")]),
    ])

    generate_tip_enums(cfg)

    proto = (cfg.proto_dir / "tip" / "alpha_error_tip.proto").read_text(
        encoding="utf-8"
    )
    assert 'option java_package = "com.game.table";' in proto


def test_empty_new_segment_does_not_claim_base_as_allocated(tmp_path: Path):
    """新空组的 Lo/Hi 占位不能让 InAllocatedRange(base) 返回 true。"""
    cfg = _make_cfg(tmp_path)
    write_tip_xlsx(cfg.tip_file, [
        ("//alpha_error base=1000 width=100", []),
    ])
    generate_tip_enums(cfg)

    seg = (Path(cfg.go.code_dir).parent / "tip" / "segments.go").read_text(encoding="utf-8")
    assert "Lo: 1000, Hi: 1000, Count: 0" in seg
    assert "s.Count > 0 && code >= s.Lo && code <= s.Hi" in seg


# ---------------------------------------------------------------- 红线 7:故障分类在表里

def test_fault_column_generates_go_fault_table(tmp_path: Path):
    """标了 fault 的码进 faults.go,没标的不进;顺序按码升序。"""
    cfg = _make_cfg(tmp_path)
    write_tip_xlsx(cfg.tip_file, [
        ("//alpha_error base=1000 width=100",
         [("AlphaReject", "拒绝"), ("AlphaRedisDown", "存储故障", 1), ("AlphaAlsoReject", "拒绝", "")]),
        ("//beta_error base=2000 width=100", [("BetaTimeout", "超时", "1")]),
    ])
    generate_tip_enums(cfg)

    src = _faults_go(cfg)
    assert "package tip" in src and "func IsFault(code uint32) bool" in src
    assert '{Code: 1001, Group: "alpha_error", Name: "AlphaRedisDown"}' in src
    assert '{Code: 2000, Group: "beta_error", Name: "BetaTimeout"}' in src
    assert "AlphaReject" not in src and "AlphaAlsoReject" not in src
    assert src.index("Code: 1001") < src.index("Code: 2000"), "Faults 按码升序"
    # faultSet 由 Faults 派生,不单独生成字面量(键宽不一时 gofmt 会重排对齐)
    assert "for _, f := range Faults" in src and "1001:" not in src


def test_fault_column_missing_header_aborts(tmp_path: Path):
    """fault 列被删/改名不能静默变成「全部不是故障」—— 那会让所有告警一起消失。"""
    cfg = _make_cfg(tmp_path)
    write_tip_xlsx(cfg.tip_file, [
        ("//alpha_error base=1000 width=100", [("A1", "")]),
    ], fault_header=None)
    with pytest.raises(TipAxisError) as exc:
        generate_tip_enums(cfg)
    assert "缺少 fault 列" in str(exc.value)

    # 表头写错名同样算缺失
    write_tip_xlsx(cfg.tip_file, [
        ("//alpha_error base=1000 width=100", [("A1", "")]),
    ], fault_header="falut")
    with pytest.raises(TipAxisError) as exc:
        generate_tip_enums(cfg)
    assert "缺少 fault 列" in str(exc.value)


def test_fault_column_is_located_by_header_not_by_position(tmp_path: Path):
    """策划在中间插一列,fault 列的含义不能跟着挪位。"""
    cfg = _make_cfg(tmp_path)
    write_tip_xlsx(cfg.tip_file, [
        ("//alpha_error base=1000 width=100", [("A1", "一", 1), ("A2", "二")]),
    ], fault_column=5)
    # C 列塞点无关内容,证明不是「第 3 列就是 fault」
    wb = openpyxl.load_workbook(cfg.tip_file)
    ws = wb.active
    ws.cell(row=1, column=3, value="note")
    ws.cell(row=_TIP_DATA_BEGIN + 2, column=3, value="这是备注,不是故障标记")
    wb.save(cfg.tip_file)

    generate_tip_enums(cfg)
    src = _faults_go(cfg)
    assert 'Name: "A1"' in src and 'Name: "A2"' not in src


@pytest.mark.parametrize("raw", [1, "1", True, "true", "yes", "是", " 1 "])
def test_fault_cell_truthy_vocabulary(tmp_path: Path, raw):
    cfg = _make_cfg(tmp_path)
    write_tip_xlsx(cfg.tip_file, [
        ("//alpha_error base=1000 width=100", [("A1", "", raw)]),
    ])
    generate_tip_enums(cfg)
    assert 'Name: "A1"' in _faults_go(cfg)


@pytest.mark.parametrize("raw", [0, "0", False, "false", "no", "否", "", "  "])
def test_fault_cell_falsy_vocabulary(tmp_path: Path, raw):
    cfg = _make_cfg(tmp_path)
    write_tip_xlsx(cfg.tip_file, [
        ("//alpha_error base=1000 width=100", [("A1", "", raw)]),
    ])
    generate_tip_enums(cfg)
    assert 'Name: "A1"' not in _faults_go(cfg)


@pytest.mark.parametrize("raw", ["maybe", "故障", "TODO", 2, 1.5, "是的"])
def test_fault_cell_unrecognized_value_aborts(tmp_path: Path, raw):
    """不猜:一个写着「TODO」的格子若被当成 truthy,就把业务拒绝刷成了告警。"""
    cfg = _make_cfg(tmp_path)
    write_tip_xlsx(cfg.tip_file, [
        ("//alpha_error base=1000 width=100", [("A1", ""), ("A2", "", raw)]),
    ])
    with pytest.raises(TipAxisError) as exc:
        generate_tip_enums(cfg)
    assert f"C{_TIP_DATA_BEGIN + 2}" in str(exc.value) and "不可识别" in str(exc.value)
    assert not (Path(cfg.go.code_dir).parent / "tip").exists(), "解析失败不能留下产物"


def test_fault_on_group_header_or_comment_row_aborts(tmp_path: Path):
    """组头/注释行标 fault 多半是行错位 —— 标记本想落在相邻的码上。不猜,直接报。"""
    cfg = _make_cfg(tmp_path)
    write_tip_xlsx(cfg.tip_file, [
        ("//alpha_error base=1000 width=100", [("A1", "")]),
    ])
    wb = openpyxl.load_workbook(cfg.tip_file)
    ws = wb.active
    ws.cell(row=_TIP_DATA_BEGIN, column=3, value=1)  # 组头行
    wb.save(cfg.tip_file)
    with pytest.raises(TipAxisError) as exc:
        generate_tip_enums(cfg)
    assert f"C{_TIP_DATA_BEGIN}" in str(exc.value) and "组头/注释" in str(exc.value)

    write_tip_xlsx(cfg.tip_file, [
        ("//alpha_error base=1000 width=100", [("A1", "")]),
    ])
    wb = openpyxl.load_workbook(cfg.tip_file)
    ws = wb.active
    ws.cell(row=_TIP_DATA_BEGIN + 2, column=1, value="// 组内说明")
    ws.cell(row=_TIP_DATA_BEGIN + 2, column=3, value=1)
    wb.save(cfg.tip_file)
    with pytest.raises(TipAxisError) as exc:
        generate_tip_enums(cfg)
    assert f"C{_TIP_DATA_BEGIN + 2}" in str(exc.value)


def test_fault_on_row_without_code_name_aborts(tmp_path: Path):
    cfg = _make_cfg(tmp_path)
    write_tip_xlsx(cfg.tip_file, [
        ("//alpha_error base=1000 width=100", [("A1", "")]),
    ])
    wb = openpyxl.load_workbook(cfg.tip_file)
    ws = wb.active
    ws.cell(row=_TIP_DATA_BEGIN + 5, column=3, value=1)  # A 列为空的行
    wb.save(cfg.tip_file)
    with pytest.raises(TipAxisError) as exc:
        generate_tip_enums(cfg)
    assert f"C{_TIP_DATA_BEGIN + 5}" in str(exc.value) and "没有码名" in str(exc.value)


def test_code_before_first_group_header_aborts(tmp_path: Path):
    """第一个组头之前的码以前被静默丢掉(没号、没 proto)—— 现在必须报。"""
    cfg = _make_cfg(tmp_path)
    write_tip_xlsx(cfg.tip_file, [
        ("Orphan", []),  # 不是组头,而是一个裸码
        ("//alpha_error base=1000 width=100", [("A1", "")]),
    ])
    with pytest.raises(TipAxisError) as exc:
        generate_tip_enums(cfg)
    assert "Orphan" in str(exc.value) and "第一个组头之前" in str(exc.value)


def test_tombstone_is_not_a_fault(tmp_path: Path):
    """删掉的码留号(墓碑)但不再有定义,也就不再有分类。"""
    cfg = _make_cfg(tmp_path)
    write_tip_xlsx(cfg.tip_file, [
        ("//alpha_error base=1000 width=100", [("Keep", "", 1), ("Gone", "", 1)]),
    ])
    generate_tip_enums(cfg)
    assert 'Name: "Gone"' in _faults_go(cfg)

    write_tip_xlsx(cfg.tip_file, [
        ("//alpha_error base=1000 width=100", [("Keep", "", 1)]),
    ])
    generate_tip_enums(cfg)
    src = _faults_go(cfg)
    assert 'Name: "Keep"' in src and 'Name: "Gone"' not in src
    assert _read_state(cfg)["groups"]["alpha_error"]["ids"]["Gone"] == 1001, "墓碑仍占号"


def test_fault_flip_does_not_touch_state(tmp_path: Path):
    """fault 是分类不是号:翻转它不需要迁移,state 一个字节都不该变。"""
    cfg = _make_cfg(tmp_path)
    write_tip_xlsx(cfg.tip_file, [
        ("//alpha_error base=1000 width=100", [("A1", ""), ("A2", "")]),
    ])
    generate_tip_enums(cfg)
    before = _state_path(cfg).read_bytes()
    assert 'Name: "A2"' not in _faults_go(cfg)

    write_tip_xlsx(cfg.tip_file, [
        ("//alpha_error base=1000 width=100", [("A1", ""), ("A2", "", 1)]),
    ])
    generate_tip_enums(cfg)
    assert 'Name: "A2"' in _faults_go(cfg)
    assert _state_path(cfg).read_bytes() == before


def test_empty_fault_table_still_generates_valid_go(tmp_path: Path):
    """一个故障码都没有是合法的(告警但不中止),产物必须仍是合法 Go。"""
    cfg = _make_cfg(tmp_path)
    write_tip_xlsx(cfg.tip_file, [
        ("//alpha_error base=1000 width=100", [("A1", "")]),
    ])
    generate_tip_enums(cfg)
    src = _faults_go(cfg)
    # gofmt 会把 `[]Fault{\n}` 改写成 `[]Fault{}`,产物要直接是后者
    assert "var Faults = []Fault{}" in src
    assert "var faultSet = func() map[uint32]struct{} {" in src
    # gofmt 要求文件以换行结尾;Jinja 默认吃掉模板末尾那个换行,生成器要补回去
    assert src.endswith("}\n")


def test_explicit_tip_ref_rejects_legacy_numeric_code(tmp_path: Path):
    cfg = _make_cfg(tmp_path)
    write_tip_xlsx(cfg.tip_file, [
        ("//alpha_error base=1000 width=100", [("Success", "成功")]),
    ])
    table_path = write_xlsx(cfg.data_dir, TableSpec(
        name="Permission",
        fields=[
            Field(name="id"),
            Field(name="result", type="repeated uint32", options="tip_ref", span=2),
        ],
        rows=[[1, 1000, 1]],
    ))
    table = read_table(table_path, cfg)
    assert table is not None

    with pytest.raises(TipAxisError) as exc:
        validate_tip_references([table], cfg)
    assert "Permission!C6=1" in str(exc.value)


def test_nested_tip_field_is_checked_without_marker(tmp_path: Path):
    cfg = _make_cfg(tmp_path)
    write_tip_xlsx(cfg.tip_file, [
        ("//alpha_error base=1000 width=100", [("Success", "成功")]),
    ])
    table_path = write_xlsx(cfg.data_dir, TableSpec(
        name="ActionState",
        fields=[
            Field(name="id"),
            Field(name="state", type="repeated { uint32 mode; uint32 tip }", span=2),
        ],
        rows=[[1, 1, 1]],
    ))
    table = read_table(table_path, cfg)
    assert table is not None

    with pytest.raises(TipAxisError) as exc:
        validate_tip_references([table], cfg)
    assert "ActionState!C6=1" in str(exc.value)


def test_tip_refs_accept_zero_and_allocated_codes(tmp_path: Path):
    cfg = _make_cfg(tmp_path)
    write_tip_xlsx(cfg.tip_file, [
        ("//alpha_error base=1000 width=100", [("Success", "成功")]),
    ])
    table_path = write_xlsx(cfg.data_dir, TableSpec(
        name="Permission",
        fields=[Field(name="id"), Field(name="result", options="tip_ref")],
        rows=[[1, 0], [2, 1000]],
    ))
    table = read_table(table_path, cfg)
    assert table is not None

    validate_tip_references([table], cfg)


def test_tip_ref_rejects_tombstoned_code(tmp_path: Path):
    """墓碑只防止号码复用，不能继续充当业务表可引用的活动码。"""
    cfg = _make_cfg(tmp_path)
    write_tip_xlsx(cfg.tip_file, [
        ("//alpha_error base=1000 width=100", [("Retired", "已退役")]),
    ])
    generate_tip_enums(cfg)
    write_tip_xlsx(cfg.tip_file, [
        ("//alpha_error base=1000 width=100", []),
    ])

    table_path = write_xlsx(cfg.data_dir, TableSpec(
        name="Permission",
        fields=[Field(name="id"), Field(name="result", options="tip_ref")],
        rows=[[1, 1000]],
    ))
    table = read_table(table_path, cfg)
    assert table is not None

    with pytest.raises(TipAxisError) as exc:
        validate_tip_references([table], cfg)
    assert "Permission!B6=1000" in str(exc.value)
