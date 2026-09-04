#!/usr/bin/env python3
"""生成 ``data/AGENTS.md`` 里的配置表索引块。

索引是给**检索**用的:人或 agent 想知道「哪张表有 bit_index」「谁引用了 Reward」
「等级列叫什么」,应该 grep 这一份,而不是逐个打开 21 个 .proto 或 21 个 xlsx。

手写的散文部分在标记之外,不会被覆盖;只有
``<!-- BEGIN GENERATED: schema-index -->`` 与 ``<!-- END GENERATED -->`` 之间的内容会重写。
**改完表记得重跑本工具**,否则索引会腐坏 —— 一份腐坏的索引比没有索引更坏。

用法::

    py tools/data_table_exporter/tools/gen_schema_index.py
    py tools/data_table_exporter/tools/gen_schema_index.py --check   # CI:腐坏即失败
"""

from __future__ import annotations

import argparse
import logging
import sys
from pathlib import Path

_HERE = Path(__file__).resolve().parent
_ROOT = _HERE.parent
if str(_ROOT) not in sys.path:
    sys.path.insert(0, str(_ROOT))

import openpyxl                                                    # noqa: E402

from core.config_loader import load_config                         # noqa: E402
from core.file_utils import list_xlsx                              # noqa: E402
from core.table_source import index_schema_protos, read_all_tables  # noqa: E402

logger = logging.getLogger("gen_schema_index")

BEGIN = "<!-- BEGIN GENERATED: schema-index -->"
END = "<!-- END GENERATED -->"


def _rows_of(xlsx: Path, sheet: str, begin_row: int) -> int:
    wb = openpyxl.load_workbook(xlsx, read_only=True)
    try:
        ws = wb[sheet]
        return sum(1 for r in ws.iter_rows(min_row=begin_row, values_only=True)
                   if r and r[0] is not None)
    finally:
        wb.close()


def _cell(items) -> str:
    return "`" + "` `".join(items) + "`" if items else "—"


def build_block(cfg) -> str:
    tables = sorted(read_all_tables(cfg), key=lambda t: t.name)
    index = index_schema_protos(cfg.schema_dir)
    by_sheet = {t.name: t for t in tables}
    rows_cache = {}
    for xlsx in list_xlsx(cfg.data_dir):
        wb = openpyxl.load_workbook(xlsx, read_only=True)
        sheet = wb.sheetnames[0] if wb.sheetnames else ""
        wb.close()
        if sheet in by_sheet:
            rows_cache[sheet] = (xlsx.name, _rows_of(xlsx, sheet, cfg.data_begin_row))

    out: list[str] = [BEGIN, ""]
    out.append("<!-- 由 tools/data_table_exporter/tools/gen_schema_index.py 生成，不要手改 -->")
    out.append("")
    out.append("| 表 | 源表 | 权威 schema | 列 | 字段 | 行 | 唯一键 | 多值键 | 索引 | 外键 | 位序 | 表达式 | 子消息 |")
    out.append("|---|---|---|---:|---:|---:|---|---|---|---|---|---|---|")
    for t in tables:
        src, nrows = rows_cache.get(t.name, ("?", 0))
        proto = index[t.name].name if t.name in index else "—"
        uniq = [c.name for c in t.table_keys if not c.is_multi_key]
        multi = [c.name for c in t.table_keys if c.is_multi_key]
        idx = [c.name for c in t.index_columns]
        fks = ["%s→%s.%s" % (c.name, c.foreign_key.target_table, c.foreign_key.target_column)
               for c in t.foreign_key_columns]
        fks += ["%s→%s.id(组)" % (c.name, c.group_foreign_key.target_table)
                for c in t.group_foreign_key_columns]
        bits = [c.name for c in t.bit_index_columns]
        exprs = ["%s:%s" % (c.name, c.expression_type) for c in t.expression_columns]
        subs = sorted(t.groups)
        out.append("| **%s** | `%s` | `%s` | %d | %d | %d | %s | %s | %s | %s | %s | %s | %s |"
                   % (t.name, src, proto, len(t.columns), len(t.field_numbers), nrows,
                      _cell(uniq), _cell(multi), _cell(idx), _cell(fks),
                      _cell(bits), _cell(exprs), _cell(subs)))

    out.append("")
    out.append("合计 **%d** 张表、**%d** 个物理列、**%d** 个进产物的字段。"
               % (len(tables), sum(len(t.columns) for t in tables),
                  sum(len(t.field_numbers) for t in tables)))
    out.append("")
    out.append(END)
    return "\n".join(out)


def main(argv: list[str] | None = None) -> int:
    ap = argparse.ArgumentParser(description="生成 data/AGENTS.md 的配置表索引块")
    ap.add_argument("config", nargs="?", default=None)
    ap.add_argument("--target", default=None, help="缺省 <data_dir>/AGENTS.md")
    ap.add_argument("--check", action="store_true", help="只检查是否最新,不写盘")
    args = ap.parse_args(argv)

    logging.basicConfig(level=logging.INFO, format="%(levelname)s %(message)s")
    cfg = load_config(args.config)
    target = Path(args.target) if args.target else (cfg.data_dir / "AGENTS.md")
    if not target.exists():
        logger.error("%s 不存在;索引块要嵌在手写知识库里,先建那份文件", target)
        return 1

    text = target.read_text(encoding="utf-8")
    if BEGIN not in text or END not in text:
        logger.error("%s 里找不到 %s / %s 标记", target.name, BEGIN, END)
        return 1

    head, rest = text.split(BEGIN, 1)
    _old, tail = rest.split(END, 1)
    block = build_block(cfg)
    new = head + block + tail

    if args.check:
        if new != text:
            logger.error("%s 的索引块已腐坏,跑一次 gen_schema_index.py 重生", target.name)
            return 1
        logger.info("%s 的索引块是最新的", target.name)
        return 0

    if new == text:
        logger.info("%s 无变化", target.name)
        return 0
    target.write_text(new, encoding="utf-8", newline="\n")
    logger.info("已更新 %s", target)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
