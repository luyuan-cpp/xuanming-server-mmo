"""Enum generators for Operator and Tip Excel files.

* **Operator.xlsx** → one ``.proto`` per group with an ``enum`` definition.
* **Tip.xlsx**      → one ``.proto`` per error group with an ``enum`` definition,
  plus a segment table and a text table.

Both use persistent ID mappings in ``state/`` so that enum values are stable
across regeneration.

tip 码轴的分配纪律（2026-09-02 起）
-----------------------------------
tip 码是**客户端可见的契约**：客户端拿 ``TipInfoMessage.id`` 去查文案。
全仓只有一条数轴，所有组共用，所以「谁发号」必须只有一个地方。

历史上这里是**扁平队尾发号**（``global_id = 所有组的最大号 + 1``），
而「段」只存在于 ``go/shared/serverbase/tipcode.go`` 里一张手抄的镜像表。
两者互不知道对方存在，后果是：往 common 组加一个码，它拿到的号是队尾的
130，落在 common 段（1-19）之外，而没有任何一步会报错。

现在改成 **段是发号器的输入**：

* 段由 ``Tip.xlsx`` 的组头行声明：``//common_error base=1000 width=1000``；
* 发号只在本组段内找空位，段满即 ``SystemExit`` 并点名是哪个组；
* ``state`` 里的号 **只增不减**：删掉 xlsx 里的一行不回收它的号，
  避免号被复用后老客户端把新含义当成旧含义；
* 生成期自检：段不重叠、码不越界、已发号不变、枚举名全局唯一。

枚举名为什么必须全局唯一:生成的 tip proto **没有 package 声明**，
protobuf 的 enum 值是 enum 的兄弟而不是子成员，于是所有 tip proto 的
枚举值处在同一个命名空间里。重名会在 protoc 阶段炸，这里提前拦。
"""

from __future__ import annotations

import json
import logging
import sys
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from jinja2 import Environment, FileSystemLoader

from core.config_loader import ExporterConfig
from core.file_utils import ensure_dirs, write_file
from core.schema import TableSchema

logger = logging.getLogger(__name__)

TIP_STATE_SCHEMA_VERSION = 2


# ===================================================================
# Operator enum generation
# ===================================================================

def generate_operator_enums(cfg: ExporterConfig) -> None:
    """Read ``Operator.xlsx`` and produce one proto file per group."""
    if not cfg.operator_file.exists():
        logger.warning("Operator file not found: %s", cfg.operator_file)
        return

    state_dir = cfg.state_dir / "operator"
    ensure_dirs(state_dir)
    id_file = state_dir / "id_pool.json"
    id_map = _load_json(id_file)

    groups = _read_operator_groups(cfg.operator_file)
    out_dir = cfg.proto_dir / "operator"
    ensure_dirs(out_dir)

    env = Environment(loader=FileSystemLoader(str(cfg.template_dir), encoding="utf-8"))
    tpl = env.get_template("operator_enum.proto.j2")

    for group_name, entries in groups.items():
        for enum_name, _ in entries:
            if enum_name not in id_map:
                id_map[enum_name] = max(id_map.values(), default=0) + 1

        content = tpl.render(
            group_name=group_name,
            entries=entries,
            id_map=id_map,
            java_package=cfg.java.package if cfg.java.enabled else "",
        )
        write_file(out_dir / f"{group_name.lower()}_operator.proto", content)
        logger.info("Generated operator proto: %s", group_name)

    _save_json(id_file, id_map)


def _read_operator_groups(file_path: Path) -> dict[str, list[tuple[str, int]]]:
    # 同 _read_tip_groups:read_only workbook 必须 close,否则在 Windows 上锁住源表。
    wb = load_workbook(file_path, read_only=True)
    try:
        ws = wb.active
        groups: dict[str, list[tuple[str, int]]] = {}
        current_group = None
        row_id = 1

        for row_idx in range(18, (ws.max_row or 18) + 1):
            cells = ws[row_idx]
            val = cells[0].value
            if val and str(val).startswith("//"):
                current_group = str(val).strip("/").strip()
                groups.setdefault(current_group, [])
            elif current_group and val:
                groups[current_group].append((str(val).strip(), row_id))
                row_id += 1

        return groups
    finally:
        wb.close()


# ===================================================================
# Tip enum generation
# ===================================================================

class TipAxisError(SystemExit):
    """段/号纪律被违反。一律中止导表，不产出半套。"""


def generate_tip_enums(cfg: ExporterConfig) -> None:
    """Read ``Tip.xlsx`` and produce per-group protos + segment/text tables."""
    tip_file = cfg.tip_file
    if not tip_file.is_file():
        raise TipAxisError(
            f"Tip.xlsx 不存在或不是文件: {tip_file}。\n"
            f"  Tip.xlsx 是 tip 码轴的唯一权威源；缺失时继续导表会部署上一批残留产物。"
        )

    state_dir = cfg.state_dir / "mapping" / "tip_enum_ids"
    ensure_dirs(state_dir)
    id_file = state_dir / "tip_enum_ids.json"
    state = _load_tip_state(id_file, bootstrap=cfg.tip_axis_bootstrap)

    groups = _read_tip_groups(tip_file)
    _assign_tip_ids(groups, state)
    _check_tip_axis(groups)

    out_dir = cfg.proto_dir / "tip"
    ensure_dirs(out_dir)

    env = Environment(loader=FileSystemLoader(str(cfg.template_dir), encoding="utf-8"))
    tpl = env.get_template("tip_enum.proto.j2")

    for group_name, g in groups.items():
        group_data = {name: g["ids"][name] for name, _text in g["entries"]}
        content = tpl.render(
            group_name=group_name,
            group_data=group_data,
            java_package=cfg.java.package if cfg.java.enabled else "",
        )
        write_file(out_dir / f"{group_name.lower()}_tip.proto", content)
        logger.info("Generated tip proto: %s (base=%d, %d codes)",
                    group_name, g["base"], len(group_data))

    _write_tip_state(id_file, groups, state)
    _generate_tip_segments(cfg, env, groups)
    _generate_tip_text(cfg, groups)


def validate_tip_references(tables: list[TableSchema], cfg: ExporterConfig) -> None:
    """在任何产物落盘前校验普通数据表里内嵌的 tip ID。

    字段名含 ``tip`` 的数值列自动纳入；语义上是 tip 码但名字看不出的字段，
    在第 4 行 options 标 ``tip_ref``。标在 repeated 字段首列时覆盖其整个 span。
    """
    tip_file = cfg.tip_file
    if not tip_file.is_file():
        raise TipAxisError(
            f"Tip.xlsx 不存在或不是文件: {tip_file}。\n"
            f"  Tip.xlsx 是 tip 码轴的唯一权威源；缺失时无法校验表内 tip 引用。"
        )

    id_file = cfg.state_dir / "mapping" / "tip_enum_ids" / "tip_enum_ids.json"
    state = _load_tip_state(id_file, bootstrap=cfg.tip_axis_bootstrap)
    groups = _read_tip_groups(tip_file)
    _assign_tip_ids(groups, state)
    _check_tip_axis(groups)
    valid_codes = {0}
    for group in groups.values():
        # state 中删过的名字必须保留为墓碑，防止 ID 被复用；但墓碑不会再生成
        # proto/text，普通表继续引用它只会得到悬空码，因此只能接受当前活动 entry。
        valid_codes.update(
            int(group["ids"][name]) for name, _text in group["entries"]
        )

    errors: list[str] = []
    checked = 0
    integer_types = {"int32", "uint32", "int64", "uint64", "sint32", "sint64"}
    for table in tables:
        explicit_fields = {
            column.name for column in table.columns if "tip_ref" in column.options
        }
        tip_columns = [
            column for column in table.columns
            if column.data_type in integer_types
            and (column.name in explicit_fields or "tip" in column.name.lower())
        ]
        if not tip_columns:
            continue

        wb = load_workbook(table.source_path, read_only=True, data_only=True)
        try:
            ws = wb[table.name]
            # 某些合法 xlsx 没写 worksheet dimension，read_only 下 max_row 会是 None；
            # iter_rows 仍能流式读到全部单元格，不能把这种表误判成“0 个引用”。
            for excel_row in ws.iter_rows(min_row=cfg.data_begin_row):
                if not excel_row:
                    continue
                row_idx = excel_row[0].row
                for column in tip_columns:
                    if column.excel_index >= len(excel_row):
                        continue
                    value = excel_row[column.excel_index].value
                    if value is None or value == "":
                        continue
                    checked += 1
                    try:
                        code = int(value)
                        if isinstance(value, float) and not value.is_integer():
                            raise ValueError
                    except (TypeError, ValueError):
                        errors.append(
                            f"{table.name}!{get_column_letter(column.excel_index + 1)}{row_idx}={value!r} 不是整数 tip ID"
                        )
                        continue
                    if code not in valid_codes:
                        errors.append(
                            f"{table.name}!{get_column_letter(column.excel_index + 1)}{row_idx}={code} "
                            f"不在当前 tip 码轴的已分配集合中"
                        )
        finally:
            wb.close()

    if errors:
        raise TipAxisError("普通数据表存在失效的 tip ID 引用:\n  " + "\n  ".join(errors))
    logger.info("Validated %d embedded tip reference(s) in data tables", checked)


# ---------------------------------------------------------------- parsing

def _parse_group_header(raw: str) -> tuple[str, int, int]:
    """``//common_error base=1000 width=1000`` → ``("common_error", 1000, 1000)``."""
    body = raw.lstrip("/").strip()
    parts = body.split()
    name = parts[0]
    kv: dict[str, str] = {}
    for tok in parts[1:]:
        if "=" not in tok:
            continue
        k, v = tok.split("=", 1)
        kv[k.strip()] = v.strip()

    if "base" not in kv or "width" not in kv:
        raise TipAxisError(
            f"tip 组头缺少段声明: '{raw}'。\n"
            f"  组头必须写成 //{name} base=<起始号> width=<段宽>。\n"
            f"  段是发号器的输入,不写就没法保证新码落在自己段内。\n"
            f"  (若这行本意是注释而非组头,请在 // 后加一个空格:'// {name} ...')"
        )
    try:
        base = int(kv["base"])
        width = int(kv["width"])
    except ValueError as exc:
        raise TipAxisError(f"tip 组头 '{raw}' 的 base/width 不是整数: {exc}") from exc
    if base <= 0 or width <= 0:
        raise TipAxisError(f"tip 组头 '{raw}' 的 base/width 必须为正数")
    return name, base, width


def _read_tip_groups(file_path: Path) -> dict[str, dict]:
    """读 Tip.xlsx。A 列码名(组头以 // 开头)，B 列中文文案。"""
    # read_only 的 workbook 在 Windows 上会一直占着文件句柄,不 close 就锁住
    # Tip.xlsx(策划下一步想打开都打不开)。原实现漏了这一步。
    wb = load_workbook(file_path, read_only=True)
    try:
        return _parse_tip_sheet(wb.active)
    finally:
        wb.close()


def _parse_tip_sheet(ws) -> dict[str, dict]:
    groups: dict[str, dict] = {}
    current = None

    for row_idx in range(18, (ws.max_row or 18) + 1):
        cells = ws[row_idx]
        val = cells[0].value
        if val is None:
            continue
        s = str(val).strip()
        if not s:
            continue
        if s.startswith("//"):
            # 组头 vs 注释靠「// 后面有没有空格」区分:
            #   //common_error base=1000 width=1000   ← 组头(紧贴)
            #   // 这是一句说明                        ← 注释(有空格)
            # 不能用「有没有 base=」来分:那样旧格式的 `//common_error`
            # 会被当成注释静默跳过,它下面的码全部落到上一组里去 ——
            # 恰好是这次要消灭的那类静默错误。
            rest = s[2:]
            if not rest or rest[0].isspace():
                continue
            name, base, width = _parse_group_header(s)
            if name in groups:
                raise TipAxisError(f"tip 组 '{name}' 在 Tip.xlsx 里出现了两次")
            groups[name] = {"base": base, "width": width, "entries": [], "ids": {}}
            current = name
        elif current:
            text = cells[1].value if len(cells) > 1 else None
            groups[current]["entries"].append((s, "" if text is None else str(text).strip()))

    if not groups:
        raise TipAxisError("Tip.xlsx 里没有解析到任何组")
    return groups


# ---------------------------------------------------------------- allocation

def _assign_tip_ids(groups: dict[str, dict], state: dict) -> None:
    """按段发号。已发过的号原样保留;新名字取本段内最小空位。"""
    state_groups = state.get("groups", {})

    # state 里的整组也是墓碑的一部分。若 Excel 把组头一并删掉/改名，
    # 只检查当前 groups 会让后来者复用旧段和旧码，等同于整组回收。
    # 要停用一个组，保留没有 entry 的空组头即可；显式迁移则必须同步改 state。
    missing_groups = sorted(set(state_groups) - set(groups))
    if missing_groups:
        formatted = ", ".join(missing_groups)
        raise TipAxisError(
            f"Tip.xlsx 缺少 state 中已有的 tip 组: {formatted}。\n"
            f"  整组删除/改名会丢掉段墓碑，使旧段和旧码可被重新分配。\n"
            f"  若只是停用，请在 Tip.xlsx 保留对应的空组头；若是迁移，"
            f"请同步迁移 state 并记录旧号映射。"
        )

    for name, g in groups.items():
        base, width = g["base"], g["width"]
        known = dict(state_groups.get(name, {}).get("ids", {}))

        # 已发号必须仍在本组段内 —— 否则说明有人改了 base 却没迁老号。
        stray = {k: v for k, v in known.items() if not (base <= int(v) < base + width)}
        if stray:
            raise TipAxisError(
                f"tip 组 '{name}' 的段是 [{base}, {base + width})，但 state 里这些已发号落在段外:\n"
                f"  {stray}\n"
                f"  说明 base 被改过而老号没迁。要么把 base 改回去，要么做一次显式重排"
                f"(并把旧号→新号写进 state 的 legacy_ids)。"
            )

        used = {int(v) for v in known.values()}
        g["ids"] = {k: int(v) for k, v in known.items()}

        cursor = base
        for enum_name, _text in g["entries"]:
            if enum_name in g["ids"]:
                continue
            while cursor in used and cursor < base + width:
                cursor += 1
            if cursor >= base + width:
                raise TipAxisError(
                    f"tip 组 '{name}' 的段 [{base}, {base + width}) 已满，"
                    f"放不下新码 '{enum_name}'。\n"
                    f"  把该组的 width 调大(注意别和下一组的 base 重叠)，或另开一组。"
                )
            g["ids"][enum_name] = cursor
            used.add(cursor)
            cursor += 1


# ---------------------------------------------------------------- checks

def _check_tip_axis(groups: dict[str, dict]) -> None:
    """段不重叠 / 码不越界 / 枚举名全局唯一。"""
    errors: list[str] = []

    # 1) 段两两不重叠
    segs = sorted(((g["base"], g["base"] + g["width"], name) for name, g in groups.items()))
    for (lo1, hi1, n1), (lo2, hi2, n2) in zip(segs, segs[1:]):
        if lo2 < hi1:
            errors.append(f"段重叠: {n1} [{lo1},{hi1}) 与 {n2} [{lo2},{hi2})")

    # 2) 每个码都在自己段内(_assign 已保证新号，这里兜住手改 state 的情况)
    for name, g in groups.items():
        lo, hi = g["base"], g["base"] + g["width"]
        for enum_name, code in g["ids"].items():
            if not (lo <= code < hi):
                errors.append(f"{name}.{enum_name}={code} 落在本组段 [{lo},{hi}) 之外")

    # 3) 枚举名全局唯一(tip proto 无 package，枚举值同处一个命名空间)
    seen: dict[str, str] = {}
    for name, g in groups.items():
        for enum_name, _text in g["entries"]:
            if enum_name in seen:
                errors.append(
                    f"枚举名 '{enum_name}' 在 {seen[enum_name]} 与 {name} 中重复。"
                    f"生成的 tip proto 无 package 声明，枚举值全局同名空间，protoc 会失败。"
                )
            seen[enum_name] = name

    # 4) 码值全局唯一(段不重叠时理应自动成立，重复即说明 state 被手改坏了)
    by_code: dict[int, str] = {}
    for name, g in groups.items():
        for enum_name, code in g["ids"].items():
            if code in by_code:
                errors.append(f"码 {code} 被 {by_code[code]} 与 {name}.{enum_name} 同时占用")
            by_code[code] = f"{name}.{enum_name}"

    if errors:
        raise TipAxisError("tip 码轴自检失败:\n  " + "\n  ".join(errors))


# ---------------------------------------------------------------- state io

def _load_tip_state(path: Path, *, bootstrap: bool = False) -> dict:
    # tip state 决定客户端可见 ID，已有文件一旦损坏绝不能退化成空状态重新发号。
    # Operator 的历史 loader 仍保持原行为；这里只为 tip 轴使用严格读取。
    try:
        with open(path, "r", encoding="utf-8") as f:
            raw = json.load(f)
    except FileNotFoundError:
        if bootstrap:
            return {"schema_version": TIP_STATE_SCHEMA_VERSION, "groups": {}, "legacy_ids": {}}
        raise TipAxisError(
            f"tip state 不存在: {path}。\n"
            f"  为防止误删 state 后从空重新发号，本次导表已中止。"
            f"只有确认该码轴从未发布过时，才可显式传 --tip-axis-bootstrap 初始化。"
        )
    except (OSError, UnicodeError, json.JSONDecodeError) as exc:
        raise TipAxisError(
            f"tip state 读取失败: {path}: {exc}。\n"
            f"  为防止静默重置并重新发号，本次导表已中止；请从版本库/备份恢复 state。"
        ) from exc

    if not isinstance(raw, dict):
        raise TipAxisError(f"tip state 格式错误: {path} 的根节点必须是 JSON object")
    ver = raw.get("schema_version")
    if ver != TIP_STATE_SCHEMA_VERSION:
        raise TipAxisError(
            f"{path} 是 v{ver or 1} 格式，本版发号器只认 v{TIP_STATE_SCHEMA_VERSION}。\n"
            f"  v1 是扁平队尾发号的产物，直接沿用会让新码落在自己段外。\n"
            f"  迁移方式见 docs/design/tip-code-axis.md（一次性重排，旧号→新号写进 legacy_ids）。"
        )
    raw.setdefault("groups", {})
    raw.setdefault("legacy_ids", {})
    if not isinstance(raw["groups"], dict) or not isinstance(raw["legacy_ids"], dict):
        raise TipAxisError(
            f"tip state 格式错误: {path} 的 groups/legacy_ids 必须是 JSON object"
        )
    return raw


def _check_tip_state_evolution(previous: dict, current: dict) -> None:
    """校验 state 相对上一版只能扩展，供 CI 对 merge-base 使用。

    单看当前快照无法知道墓碑是否曾存在；因此“只增不减”必须同时有历史比较。
    首次 v1 -> v2 迁移则要求每个旧号都完整登记进 ``legacy_ids``。
    """
    if previous.get("schema_version") != TIP_STATE_SCHEMA_VERSION:
        legacy = current.get("legacy_ids", {})
        missing: list[str] = []
        for group, ids in previous.items():
            if not isinstance(ids, dict):
                raise TipAxisError(f"上一版 tip state 的组 '{group}' 不是 JSON object")
            for name, old_code in ids.items():
                new_group = current.get("groups", {}).get(group, {})
                new_code = new_group.get("ids", {}).get(name)
                record = legacy.get(str(old_code))
                expected = {"group": group, "name": name, "new": new_code}
                if new_code is None or record != expected:
                    missing.append(f"{group}.{name}: {old_code} -> {new_code}")
        if missing:
            raise TipAxisError(
                "tip state 的 v1 -> v2 迁移记录不完整:\n  " + "\n  ".join(missing)
            )
        return

    previous_groups = previous.get("groups", {})
    current_groups = current.get("groups", {})
    errors: list[str] = []
    for group, old_group in previous_groups.items():
        new_group = current_groups.get(group)
        if new_group is None:
            errors.append(f"删除了已有组 {group}")
            continue

        old_base = int(old_group["base"])
        old_width = int(old_group["width"])
        new_base = int(new_group["base"])
        new_width = int(new_group["width"])
        if new_base > old_base or new_base + new_width < old_base + old_width:
            errors.append(
                f"缩小/平移了已有段 {group}: [{old_base},{old_base + old_width}) -> "
                f"[{new_base},{new_base + new_width})"
            )

        new_ids = new_group.get("ids", {})
        for name, old_code in old_group.get("ids", {}).items():
            if name not in new_ids:
                errors.append(f"删除了已有号 {group}.{name}={old_code}")
            elif int(new_ids[name]) != int(old_code):
                errors.append(
                    f"改写了已有号 {group}.{name}: {old_code} -> {new_ids[name]}"
                )

    previous_legacy = previous.get("legacy_ids", {})
    current_legacy = current.get("legacy_ids", {})
    for old_code, record in previous_legacy.items():
        if current_legacy.get(old_code) != record:
            errors.append(f"删除/改写了 legacy_ids[{old_code}]")

    if errors:
        raise TipAxisError("tip state 违反只增不减约束:\n  " + "\n  ".join(errors))


def _write_tip_state(path: Path, groups: dict[str, dict], state: dict) -> None:
    """写回 state。ids 只增不减:xlsx 里删掉的名字仍留在 state 中，号不回收。"""
    out_groups = dict(state.get("groups", {}))
    for name, g in groups.items():
        prev = dict(out_groups.get(name, {}).get("ids", {}))
        prev.update(g["ids"])            # 只增不减
        out_groups[name] = {
            "base": g["base"],
            "width": g["width"],
            "ids": {k: int(v) for k, v in sorted(prev.items(), key=lambda kv: int(kv[1]))},
        }
    state["schema_version"] = TIP_STATE_SCHEMA_VERSION
    state["groups"] = out_groups
    _save_json(path, state)


# ---------------------------------------------------------------- extra outputs

def _generate_tip_segments(cfg: ExporterConfig, env: Environment, groups: dict[str, dict]) -> None:
    """生成 Go 侧段表。

    取代 go/shared/serverbase/tipcode.go 里那张手抄的 tipDomains ——
    手抄的镜像是这套机制上一次失效的根因。
    """
    if not cfg.go.enabled:
        return
    rows = []
    for name, g in sorted(groups.items(), key=lambda kv: kv[1]["base"]):
        codes = sorted(g["ids"].values())
        rows.append({
            "domain": name.removesuffix("_error") or name,
            "group": name,
            "base": g["base"],
            "width": g["width"],
            "lo": codes[0] if codes else g["base"],
            "hi": codes[-1] if codes else g["base"],
            "count": len(codes),
        })
    out_dir = Path(cfg.go.code_dir).parent / "tip"
    ensure_dirs(out_dir)
    tpl = env.get_template("tip_segments.go.j2")
    write_file(out_dir / "segments.go", tpl.render(rows=rows))
    logger.info("Generated tip segment table: %d segments", len(rows))


def _generate_tip_text(cfg: ExporterConfig, groups: dict[str, dict]) -> None:
    """生成 id → 文案 的产物。

    Tip.xlsx 的 B 列(design 拥有的中文文案)在此之前**没有任何出口** ——
    只有 A 列被读来生成枚举，于是「客户端按 id 查文案」这条链一直是断的。
    """
    payload: dict[str, str] = {}
    missing: list[str] = []
    for name, g in groups.items():
        for enum_name, text in g["entries"]:
            code = g["ids"][enum_name]
            if text:
                payload[str(code)] = text
            else:
                missing.append(f"{name}.{enum_name}({code})")

    ensure_dirs(cfg.json_dir)
    out = Path(cfg.json_dir) / "tip_text.json"
    write_file(out, json.dumps(payload, indent=2, ensure_ascii=False, sort_keys=True))
    logger.info("Generated tip text table: %d/%d codes have text",
                len(payload), len(payload) + len(missing))
    if missing:
        logger.warning("以下 tip 码没有中文文案(客户端只能按既有 fallback 处理): %d 个", len(missing))
        for m in missing[:20]:
            logger.warning("  缺文案: %s", m)
        if len(missing) > 20:
            logger.warning("  ...另有 %d 个", len(missing) - 20)


# ===================================================================
# Shared helpers
# ===================================================================

def _load_json(path: Path) -> dict:
    if path.exists():
        try:
            with open(path, "r", encoding="utf-8") as f:
                return json.load(f)
        except Exception as exc:
            logger.error("Failed to load %s: %s", path, exc)
    return {}


def _save_json(path: Path, data: dict) -> None:
    with open(path, "w", encoding="utf-8") as f:
        json.dump(data, f, indent=2, ensure_ascii=False)
