#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""装备属性系统(docs/design/equipment-attributes.md §2)要对 xlsx 做的全部改动,写成可重复执行的脚本。

为什么是脚本而不是手改:
  - data/Item.xlsx、data/EquipSlot.xlsx、data/tip/Tip.xlsx、data/MessageLimiter.xlsx 同时有多个会话在写,
    它们是二进制、不能 3-way 合并,整表写回会**静默**删掉别人的行(导表器照常跑通,只是少几行 / 少发几个号)。
    唯一安全的写法是 openpyxl 从磁盘重新载入 → 只往**空格子**里写自己那几列 / 那几行 → save → 读回来逐格核对。
    本脚本对既有表的每一次写入都先断言目标格为空,所以「别人的格子一格不动」是构造保证,不是约定。
  - 导表器要求 schema 与 xlsx **成对**:data/schema/ 下已经有 4 张新表的 schema,并且给 Item / EquipSlot 加了字段。
    只有 schema 没有对应的 xlsx(或列)时,任何人跑 `dev.bat gen` 都会整批失败、产物一个字节不落盘。
    所以 new-tables / item-columns / equip-slots 必须先于下一次导表执行,并与 schema 同批提交。
    新建的 xlsx 只在一台机器上生成并提交(openpyxl 每次写出的字节都不同,两台机器各生成一份会在 git 里
    撞成二进制 add/add 冲突),其它机器等 pull。
  - MessageLimiter 按**消息号**登记档位,而新方法落到哪个号要等全量 proto-gen 之后才定,
    所以 message-limiter 按方法名去 proto/message_id.txt 查号,不把号写死。

六个子命令都**幂等**:已经是目标状态就只报告、不改文件;发现与预期不符的现状就中止,绝不覆盖。

用法(仓库根目录执行;需要 Python 3 + openpyxl + pyyaml,与导表器同一套依赖):

    python tools/scripts/equip_xlsx_patch.py new-tables        # ① 建 EquipAttribute / EquipAttributeCap / EquipAffixPool / EquipAffixRule
    python tools/scripts/equip_xlsx_patch.py item-columns      # ② Item.xlsx 追加 8 个字段(13 个物理列)
    python tools/scripts/equip_xlsx_patch.py item-rows         # ③ Item.xlsx 追加 20 件样例装备(依赖 ① ②)
    python tools/scripts/equip_xlsx_patch.py equip-slots       # ④ EquipSlot.xlsx 加 name 列 + 槽 3–6
    python tools/scripts/equip_xlsx_patch.py tip-codes         # ⑤ Tip.xlsx 新开 //equip_error 段 + 11 个码
    #   —— 第一次导表(dev.bat gen)→ 全量 proto-gen ——
    python tools/scripts/equip_xlsx_patch.py message-limiter   # ⑥ 全量 proto-gen 之后、第二次导表之前

在哪棵树上跑导表 / proto-gen(2026-10-08 整体审查补记,顺序不对会留下查不出来的错):
  - 「第一次导表 → 全量 proto-gen」**不要在没并入 origin/main 的隔离工作树上原样跑**:
      · 发号:基线里有别人已提交、尚未 regen 的 RPC(帮会 5 个等),而生成器给无号方法发号是从 Go map 里取,
        顺序随机 —— 先跑的一方会把 origin/main 已经发出去的 239–243 发给别的方法。先让这棵树带上
        origin/main 的 proto/message_id.txt(整体并入 main),再 proto-gen;跑完核对 main 已有的号逐行不变、
        新号只出现在其后。
      · 客户端落点:proto_gen.yaml 的 unity_client_dir 与 exporter_config.yaml 的 csharp.deploy 都是相对路径
        ../mmorpg-client,从别名工作树跑会写进**另一个会话**的客户端检出。用配置副本改指向(或关掉)再跑。
      · 子模块:隔离工作树的 third_party/ 是空的,`dev.bat gen` 第一步找不到 protoc 就失败;C++ 也只能在
        子模块齐全的检出里编译。
  - ⑥ message-limiter 按**数字消息号**往二进制的 MessageLimiter.xlsx 里写行:只在「最终要进 main 的那棵树」上、
    消息号已经定死之后跑。号一旦重发,先前写下的行会留在旧号上,变成给别的 RPC 限流(脚本只认得出
    「那个号上已有一行不同的档位」,认不出「这一行原本是我写的」)。
  完整的先后顺序、命令与通过标准见 docs/design/equipment-attributes.md §9。

加 --dry-run 只打印将要做什么,不写文件。执行前先确认
`git status --short data/Item.xlsx data/EquipSlot.xlsx data/tip/Tip.xlsx data/MessageLimiter.xlsx`
没有别人的未提交改动(任一时刻只允许一方有未提交修改)。

数值出处:docs/design/equipment-attributes.md §2.1–§2.7(全部是占位建议,D9)。改数值 = 改本文件顶部的常量表
再重新生成;表一旦被策划手调过,new-tables 会拒绝覆盖,此后直接改 xlsx。
"""

from __future__ import annotations

import argparse
import json
import sys
from dataclasses import dataclass
from pathlib import Path

try:
    import openpyxl
    from openpyxl.utils import get_column_letter
except ImportError:  # pragma: no cover - 环境问题,给一句能照着做的话
    sys.exit("缺少 openpyxl:请先执行  pip install openpyxl")

REPO_ROOT = Path(__file__).resolve().parents[2]
DATA_DIR = REPO_ROOT / "data"
SCHEMA_DIR = DATA_DIR / "schema"
EXPORTER_DIR = REPO_ROOT / "tools" / "data_table_exporter"

# 列名、列宽、第 5 行中文说明都不在这里写第二份:直接用导表器自己的 schema 解析器读权威 schema,
# 写完再用导表器自己的绑定器(build_table_schema)核对「schema 与 xlsx 对得上」。
# 这样本脚本说"好了"与导表器说"好了"是同一个判据,不会出现脚本通过而导表失败。
if str(EXPORTER_DIR) not in sys.path:
    sys.path.insert(0, str(EXPORTER_DIR))
try:
    from core.schema_proto import SchemaProtoError, build_table_schema, parse_schema_proto
except ImportError as exc:  # pragma: no cover - 环境问题
    sys.exit(f"无法载入导表器的 schema 解析器({exc}):请先执行  pip install pyyaml openpyxl")

# 主表版式(data/AGENTS.md):第 1 行列名 / 第 2-4 行留空 / 第 5 行中文说明(schema 注释首行的投影)/ 第 6 行起数据。
HEADER_ROW = 1
COMMENT_ROW = 5
DATA_BEGIN_ROW = 6


# ════════════════════════════════════════════════════════════════════════════
# 数据常量表(设计文档 §2.1–§2.7)。改数值只改这一段。
# ════════════════════════════════════════════════════════════════════════════

# §2.1 属性清单。id 一经使用不再改义;C++ 只按 effect 分支。
#   effect:1 一级属性点(param = AttributeDimension.id)/ 2 所有一级属性点 / 3 二级属性平加
#          (param:1 气血上限 2 法力上限 3 物伤 4 法伤 5 速度 6 防御)/ 4 伤害(物伤与法伤各加)/
#          5 战斗类属性(param = CombatAttributes 字段号)
#   percent:0 点数 / 1 整数百分点
EQUIP_ATTRIBUTES = [
    # id  显示名              effect  param  percent
    (1,  "伤害",               4,      0,     0),
    (2,  "准确",               3,      3,     0),
    (3,  "力量",               1,      103,   0),
    (4,  "体质",               1,      101,   0),
    (5,  "灵力",               1,      102,   0),
    (6,  "敏捷",               1,      104,   0),
    (7,  "所有属性",           2,      0,     0),
    (8,  "防御",               3,      6,     0),
    (9,  "气血",               3,      1,     0),
    (10, "法力",               3,      2,     0),
    (11, "速度",               3,      5,     0),   # 不进任何池,只作鞋子的基础属性
    (12, "物理必杀率",         5,      1,     1),
    (13, "法术必杀率",         5,      2,     1),
    (14, "物理连击率",         5,      3,     1),
    (15, "反击率",             5,      4,     1),
    (16, "反震率",             5,      5,     1),
    (17, "所有技能上升",       5,      6,     0),
    (18, "忽视所有抗异常",     5,      7,     1),
    (19, "抗中毒",             5,      8,     1),
    (20, "抗冰冻",             5,      9,     1),
    (21, "抗昏睡",             5,      10,    1),
    (22, "抗遗忘",             5,      11,    1),
    (23, "抗混乱",             5,      12,    1),
    (24, "所有抗异常",         5,      13,    1),
    (25, "抗法术",             5,      14,    1),
    (26, "抗物理",             5,      15,    1),
]
# 显示排序:设计文档没给初值,落地为 id × 10(升序即上表顺序,留空档方便策划插队)。
EQUIP_ATTRIBUTE_SORT_STEP = 10

# §2.6 五个装备等级档;样例装备与上限表共用。
EQUIP_LEVEL_TIERS = (1, 20, 40, 60, 80)

# §2.6 上限(EquipAttributeCap.cap),每行五个数对应上面五个等级档。
# min_value = max(1, ceil(cap × 0.3));行 id = attr_id * 1000 + level。
EQUIP_ATTRIBUTE_CAPS = [
    # 属性(同一行的属性共用一组上限)                                          L1   L20   L40   L60   L80
    (("伤害",),                                                               (20,  300,  600,  900,  1200)),
    (("准确",),                                                               (30,  400,  800,  1200, 1600)),
    (("力量", "体质", "灵力", "敏捷"),                                        (2,   5,    10,   15,   20)),
    (("所有属性",),                                                           (1,   4,    8,    12,   16)),
    (("防御",),                                                               (6,   80,   160,  240,  320)),
    (("气血",),                                                               (10,  160,  320,  480,  640)),
    (("法力",),                                                               (8,   120,  240,  360,  480)),
    (("物理必杀率", "法术必杀率", "物理连击率", "反击率", "反震率"),          (2,   4,    6,    8,    10)),
    (("所有技能上升",),                                                       (1,   2,    3,    4,    5)),
    (("忽视所有抗异常",),                                                     (2,   5,    10,   15,   20)),
    (("抗中毒", "抗冰冻", "抗昏睡", "抗遗忘", "抗混乱"),                      (2,   5,    10,   15,   20)),
    (("所有抗异常",),                                                         (2,   4,    8,    12,   15)),
    (("抗法术", "抗物理"),                                                    (1,   2,    3,    4,    5)),
]

# §2.3 属性池(顺序 = 设计文档 §1.1 的用户需求原文)。行 id = pool_id * 1000 + attr_id。
AFFIX_POOL_WEAPON = 1
AFFIX_POOL_ARMOR = 2
EQUIP_AFFIX_POOLS = {
    AFFIX_POOL_WEAPON: ("伤害", "力量", "体质", "灵力", "敏捷", "准确", "物理连击率", "反击率",
                        "物理必杀率", "所有技能上升", "忽视所有抗异常", "法术必杀率"),
    AFFIX_POOL_ARMOR: ("防御", "气血", "法力", "力量", "体质", "灵力", "敏捷", "所有属性", "反震率",
                       "抗中毒", "抗冰冻", "抗昏睡", "抗遗忘", "抗混乱", "所有抗异常", "抗法术", "抗物理"),
}
EQUIP_AFFIX_POOL_SIZES = {AFFIX_POOL_WEAPON: 12, AFFIX_POOL_ARMOR: 17}   # 用户需求给的条数,防漏抄
EQUIP_AFFIX_WEIGHT = 100

# §2.4 掷值规则:蓝 1–3 条,粉 30%,黄 15%(万分比)。
EQUIP_AFFIX_RULES = [
    dict(id=1, blue_min=1, blue_max=3, pink_rate=3000, yellow_rate=1500),
]
SAMPLE_AFFIX_RULE = 1

# §2.6 部位与槽位(EquipSlot 追加行)。装备栏容量是 10(bag_system.h 的 kEquipmentCapacity),槽号必须 < 10。
# 既有 0–2 号槽与部位 1 / 2 是单测夹具,不动。
EQUIP_BAG_CAPACITY = 10
EQUIP_SLOTS = [
    # 槽 id  equip_kind  name
    (3,      11,         "武器"),
    (4,      12,         "帽子"),
    (5,      13,         "衣服"),
    (6,      14,         "鞋子"),
]

# §2.6 样例装备:4 个部位 × 5 个等级档 = 20 件。
#   id = 1000 + 部位序 * 100 + 等级档序(1..5);基础属性值 = 每级系数 × equip_level + 常数。
#   图标:武器取自客户端仓 Assets/Resources/UI/qdao_v3/icons_weapon/(不带扩展名);防具暂无图标,留空走客户端兜底。
WEAPON_ICON_DIR = "UI/qdao_v3/icons_weapon/"
SAMPLE_EQUIP_PARTS = [
    dict(order=1, equip_kind=11, affix_pool=AFFIX_POOL_WEAPON,
         names=("新手木剑", "青锋剑", "寒铁剑", "流云剑", "惊雷剑"),
         icons=("012_peachwood_sword", "007_flying_sword", "001_daoist_saber",
                "074_lu_dongbin_immortal_sword", "19_lu_dongbin_immortal_sword"),
         description="{level}级武器,提供伤害。",
         base_attr=(("伤害", 20, 20),)),                       # 伤害 = 20L + 20
    dict(order=2, equip_kind=12, affix_pool=AFFIX_POOL_ARMOR,
         names=("布头巾", "青纱冠", "寒铁盔", "流云冠", "惊雷冠"),
         icons=None,
         description="{level}级帽子,提供防御。",
         base_attr=(("防御", 6, 6),)),                         # 防御 = 6L + 6
    dict(order=3, equip_kind=13, affix_pool=AFFIX_POOL_ARMOR,
         names=("粗布衣", "青纱袍", "寒铁甲", "流云袍", "惊雷甲"),
         icons=None,
         description="{level}级衣服,提供防御。",
         base_attr=(("防御", 12, 12),)),                       # 防御 = 12L + 12
    dict(order=4, equip_kind=14, affix_pool=AFFIX_POOL_ARMOR,
         names=("草鞋", "青纱履", "寒铁靴", "流云履", "惊雷靴"),
         icons=None,
         description="{level}级鞋子,提供防御与速度。",
         base_attr=(("防御", 6, 6), ("速度", 4, 8))),          # 防御 = 6L + 6、速度 = 4L + 8
]

# Item.xlsx 新列的追加顺序。**最右一列必须是标量**:导表器用 max_column 推最后一个字段的列宽,
# 数组 / 子消息列排在最右时,只要没有任何一行把最后一槽填满,列宽就对不上 (cfg_slots)。
ITEM_NEW_COLUMNS = ("name", "description", "icon_key", "equip_level", "equip_class", "affix_pool",
                    "base_attr", "affix_rule")

# §2.7 tip 码(A 列不带 k;文案标点随设计文档用半角,与 pet_error 段一致)。fault=True 的是服务端内部故障。
TIP_EQUIP_GROUP = "equip_error"
TIP_EQUIP_BASE = 28000
TIP_EQUIP_WIDTH = 1000
TIP_EQUIP_CODES = [
    # 码名                    文案                              fault
    ("EquipItemNotFound",     "物品不存在",                     False),
    ("EquipNotEquipment",     "该物品不是装备",                 False),
    ("EquipLevelNotEnough",   "等级不足,无法装备",              False),
    ("EquipClassMismatch",    "职业不符,无法装备",              False),
    ("EquipNoSlot",           "没有可用的装备栏位",             False),
    ("EquipNotEquipped",      "该装备未穿戴",                   False),
    ("EquipBagFull",          "背包已满,无法卸下",              False),
    ("EquipInBattle",         "战斗中无法更换装备",             False),
    ("EquipFrozen",           "当前状态无法更换装备",           False),
    ("EquipGrantInvalid",     "发放参数无效",                   False),
    ("EquipInternalError",    "装备操作失败,请稍后再试",        True),
]

# MessageLimiter:读 10 次/秒、写 5 次/秒,与表内既有行同档;time_window / tip_message 照抄既有行
# (1 秒 / 1000 = kRateLimitExceeded)。
LIMITER_SERVICE = "SceneBagClientPlayer"
LIMITER_NEW_METHODS = {          # 本批新增的方法:号必须查得到;表里已有该号时档位必须一致
    "EquipItem": 5,
    "UnequipItem": 5,
    "GmGrantItem": 5,
}
LIMITER_BACKFILL_METHODS = {     # 早就有的方法:不在表里才补;已在表里就尊重现有档位
    "GetBag": 10,
    "SortBag": 5,
}
LIMITER_TIME_WINDOW = 1
LIMITER_TIP_MESSAGE = 1000


# ════════════════════════════════════════════════════════════════════════════
# 通用工具
# ════════════════════════════════════════════════════════════════════════════

def fail(message: str) -> "NoReturn":
    sys.exit(f"[中止] {message}")


def rel(path: Path) -> str:
    return path.relative_to(REPO_ROOT).as_posix()


def only_sheet(wb, path: Path, expect_title: str | None = None):
    if len(wb.sheetnames) != 1:
        fail(f"{path.name} 预期只有一个 sheet,实际 {wb.sheetnames}")
    ws = wb[wb.sheetnames[0]]
    if expect_title is not None and ws.title != expect_title:
        fail(f"{path.name} 的 sheet 名应为 {expect_title!r}(导表器按它匹配 schema),实际 {ws.title!r}")
    return ws


def sheet_matrix(ws) -> list[list]:
    return [[ws.cell(row=r, column=c).value for c in range(1, ws.max_column + 1)]
            for r in range(1, ws.max_row + 1)]


def _trimmed(matrix: list[list]) -> dict[tuple[int, int], object]:
    """矩阵 -> {(行, 列): 值},只留非空格。两张表「非空格完全相同」即内容相同,与尾部留白多少无关。"""
    return {(r, c): v for r, row in enumerate(matrix, start=1) for c, v in enumerate(row, start=1)
            if v is not None}


def first_difference(actual: list[list], expected: list[list]) -> str | None:
    a, e = _trimmed(actual), _trimmed(expected)
    if a == e:
        return None
    r, c = min(k for k in set(a) | set(e) if a.get(k) != e.get(k))
    return f"{get_column_letter(c)}{r}:读回 {a.get((r, c))!r},预期 {e.get((r, c))!r}"


def header_cells(ws) -> dict[str, int]:
    """第 1 行非空表头 -> {列名: 1 起的列号}。重名直接中止(导表器也会拒绝)。"""
    out: dict[str, int] = {}
    for c in range(1, ws.max_column + 1):
        value = ws.cell(row=HEADER_ROW, column=c).value
        name = str(value).strip() if value is not None else ""
        if not name:
            continue
        if name in out:
            fail(f"{ws.title} 第 1 行有重复表头 {name!r}")
        out[name] = c
    return out


Patch = tuple[int, int, object]     # (行, 列, 值),行列都从 1 起


def apply_patches_and_verify(xlsx_path: Path, sheet: str | None, wb, ws, before: list[list],
                             patches: list[Patch], bind_schema: bool) -> None:
    """只往空格子里写 → save → 读回逐格核对 →(主表)用导表器的绑定器再核一遍。

    「目标格必须为空」是本脚本对既有表的全部写入纪律:它保证别人的格子一格都不会被改,
    也保证重复执行不会把同一批数据叠写两遍。
    """
    for row, col, value in patches:
        if value is None:
            fail(f"内部错误:补丁里出现了空值({get_column_letter(col)}{row})")
        if ws.cell(row=row, column=col).value is not None:
            fail(f"{xlsx_path.name} 的 {get_column_letter(col)}{row} 已有内容 "
                 f"{ws.cell(row=row, column=col).value!r},本脚本只写空格、不覆盖 —— 先查清是谁写的。")
    for row, col, value in patches:
        ws.cell(row=row, column=col, value=value)
    wb.save(xlsx_path)

    width = max([len(r) for r in before] + [c for _r, c, _v in patches])
    height = max([len(before)] + [r for r, _c, _v in patches])
    expected = [list(r) + [None] * (width - len(r)) for r in before]
    expected += [[None] * width for _ in range(height - len(expected))]
    for row, col, value in patches:
        expected[row - 1][col - 1] = value

    check_ws = only_sheet(openpyxl.load_workbook(xlsx_path), xlsx_path, sheet)
    diff = first_difference(sheet_matrix(check_ws), expected)
    if diff:
        fail(f"保存后读回的 {xlsx_path.name} 与预期不符({diff})—— 立刻 git checkout -- {rel(xlsx_path)} 还原")
    if bind_schema:
        bind(sheet, xlsx_path)


# ── 权威 schema ──────────────────────────────────────────────────────────────

@dataclass(frozen=True)
class FieldLayout:
    name: str                       # 第 1 行表头名
    comment: str                    # schema 注释首行,投影到第 5 行
    width: int                      # 占几个物理列
    sub_names: tuple[str, ...]      # 子消息数组每条的子列名(展开后的列名);标量为空


def schema_path(sheet: str) -> Path:
    return SCHEMA_DIR / f"{sheet.lower()}_table.proto"


def schema_layout(sheet: str) -> list[FieldLayout]:
    """从权威 schema 读出每个字段的表头名、注释首行与列宽。顺序 = schema 声明顺序。"""
    path = schema_path(sheet)
    if not path.exists():
        fail(f"找不到权威 schema {rel(path)}")
    try:
        messages = parse_schema_proto(path)
    except SchemaProtoError as exc:
        fail(f"{rel(path)} 解析失败:{exc}")
    mains = [m for m in messages.values() if m.opts.get("cfg_sheet") == sheet]
    if len(mains) != 1:
        fail(f"{rel(path)} 里应当恰好有一个 (cfg_sheet) = {sheet!r} 的 message")
    layout: list[FieldLayout] = []
    for field in mains[0].fields:
        if field.opts.get("cfg_col") or field.opts.get("cfg_owner"):
            fail(f"{path.name} 的字段 {field.name} 用了 (cfg_col) / (cfg_owner),本脚本没处理这种列,先同步脚本")
        comment = field.comment.splitlines()[0].strip() if field.comment else ""
        slots = field.opts.get("cfg_slots")
        if field.type_expr in messages:                         # 子消息数组:每条占 len(子字段) 列
            subs = sorted(messages[field.type_expr].fields, key=lambda f: f.number)
            if not slots:
                fail(f"{path.name} 的字段 {field.name} 是子消息数组,必须写 (cfg_slots)")
            layout.append(FieldLayout(field.name, comment, slots * len(subs), tuple(f.name for f in subs)))
        elif field.repeated or field.type_expr.startswith("map"):
            fail(f"{path.name} 的字段 {field.name} 是标量数组 / map,本脚本没处理这种列,先同步脚本")
        else:
            layout.append(FieldLayout(field.name, comment, 1, ()))
    if not layout or layout[0].name != "id":
        fail(f"{path.name} 的首个字段应为 id")
    return layout


def bind(sheet: str, xlsx_path: Path):
    """用导表器自己的绑定器核对 schema 与 xlsx;返回 TableSchema。"""
    try:
        return build_table_schema(schema_path(sheet), xlsx_path)
    except SchemaProtoError as exc:
        fail(f"权威 schema 与 {rel(xlsx_path)} 对不上(导表也会因此失败):{exc}")


# ── 常量表自检 ───────────────────────────────────────────────────────────────

ATTR_ID_BY_NAME = {name: attr_id for attr_id, name, _e, _p, _pc in EQUIP_ATTRIBUTES}


def attr_id(name: str) -> int:
    if name not in ATTR_ID_BY_NAME:
        fail(f"常量表引用了不存在的属性名 {name!r}")
    return ATTR_ID_BY_NAME[name]


def self_check() -> None:
    """常量表之间的引用关系。任何一条不成立都说明抄错了,不带着错数据去写表。"""
    ids = [row[0] for row in EQUIP_ATTRIBUTES]
    names = [row[1] for row in EQUIP_ATTRIBUTES]
    if len(set(ids)) != len(ids) or len(set(names)) != len(names):
        fail("EQUIP_ATTRIBUTES 的 id 或显示名有重复")

    pooled: set[str] = set()
    for pool_id, members in EQUIP_AFFIX_POOLS.items():
        if len(members) != EQUIP_AFFIX_POOL_SIZES[pool_id] or len(set(members)) != len(members):
            fail(f"属性池 {pool_id} 应有 {EQUIP_AFFIX_POOL_SIZES[pool_id]} 条且互不重复,实际 {len(members)} 条")
        for name in members:
            attr_id(name)
        pooled.update(members)

    capped = [name for group, _caps in EQUIP_ATTRIBUTE_CAPS for name in group]
    if len(set(capped)) != len(capped):
        fail("EQUIP_ATTRIBUTE_CAPS 里有属性出现了两次")
    # 进池的属性必须有上限(否则掷不出来),有上限的属性必须进池(否则是死数据)。
    if set(capped) != pooled:
        fail(f"上限表与属性池不一致:只在上限表 {sorted(set(capped) - pooled)},只在池里 {sorted(pooled - set(capped))}")
    for group, caps in EQUIP_ATTRIBUTE_CAPS:
        if len(caps) != len(EQUIP_LEVEL_TIERS) or any(c <= 0 for c in caps) or list(caps) != sorted(caps):
            fail(f"{group} 的上限应为 {len(EQUIP_LEVEL_TIERS)} 个不减的正数,实际 {caps}")
    if max(EQUIP_LEVEL_TIERS) >= 1000 or max(ids) >= 1000:
        fail("行号公式 attr_id * 1000 + level / pool_id * 1000 + attr_id 要求等级与属性 id 都 < 1000")

    slot_ids = [s for s, _k, _n in EQUIP_SLOTS]
    if len(set(slot_ids)) != len(slot_ids) or any(s >= EQUIP_BAG_CAPACITY for s in slot_ids):
        fail(f"EQUIP_SLOTS 的槽号必须互不重复且 < 装备栏容量 {EQUIP_BAG_CAPACITY}")
    slot_kinds = {k for _s, k, _n in EQUIP_SLOTS}
    rule_ids = {r["id"] for r in EQUIP_AFFIX_RULES}
    if SAMPLE_AFFIX_RULE not in rule_ids:
        fail(f"样例装备引用的掷值规则 {SAMPLE_AFFIX_RULE} 不在 EQUIP_AFFIX_RULES 里")
    for part in SAMPLE_EQUIP_PARTS:
        if part["equip_kind"] not in slot_kinds:
            fail(f"样例装备的部位 {part['equip_kind']} 没有任何槽位接受")
        if part["affix_pool"] not in EQUIP_AFFIX_POOLS:
            fail(f"样例装备引用了不存在的属性池 {part['affix_pool']}")
        if len(part["names"]) != len(EQUIP_LEVEL_TIERS):
            fail(f"部位 {part['equip_kind']} 的名称应有 {len(EQUIP_LEVEL_TIERS)} 个")
        if part["icons"] is not None and len(part["icons"]) != len(EQUIP_LEVEL_TIERS):
            fail(f"部位 {part['equip_kind']} 的图标应有 {len(EQUIP_LEVEL_TIERS)} 个")
        for name, _per_level, _const in part["base_attr"]:
            attr_id(name)

    tip_names = [name for name, _t, _f in TIP_EQUIP_CODES]
    if len(set(tip_names)) != len(tip_names) or len(tip_names) > TIP_EQUIP_WIDTH:
        fail("TIP_EQUIP_CODES 码名重复或超过段宽")


# ════════════════════════════════════════════════════════════════════════════
# ① new-tables
# ════════════════════════════════════════════════════════════════════════════

def equip_attribute_rows() -> list[dict]:
    return [dict(id=i, name=name, effect=effect, effect_param=param, percent=percent,
                 sort=i * EQUIP_ATTRIBUTE_SORT_STEP)
            for i, name, effect, param, percent in EQUIP_ATTRIBUTES]


def equip_attribute_cap_rows() -> list[dict]:
    rows = []
    for group, caps in EQUIP_ATTRIBUTE_CAPS:
        for name in group:
            for level, cap in zip(EQUIP_LEVEL_TIERS, caps):
                # ceil(cap × 0.3) 用整数算:浮点的 0.3 不是精确值,乘出来再 ceil 可能多进一位。
                min_value = max(1, (cap * 3 + 9) // 10)
                rows.append(dict(id=attr_id(name) * 1000 + level, attr_id=attr_id(name), level=level,
                                 min_value=min_value, cap=cap))
    return sorted(rows, key=lambda r: r["id"])


def equip_affix_pool_rows() -> list[dict]:
    rows = [dict(id=pool_id * 1000 + attr_id(name), pool_id=pool_id, attr_id=attr_id(name),
                 weight=EQUIP_AFFIX_WEIGHT)
            for pool_id, members in EQUIP_AFFIX_POOLS.items() for name in members]
    return sorted(rows, key=lambda r: r["id"])


def new_tables() -> dict[str, list[dict]]:
    return {
        "EquipAttribute": equip_attribute_rows(),
        "EquipAttributeCap": equip_attribute_cap_rows(),
        "EquipAffixPool": equip_affix_pool_rows(),
        "EquipAffixRule": [dict(r) for r in EQUIP_AFFIX_RULES],
    }


def cmd_new_tables(dry_run: bool) -> None:
    for sheet, rows in new_tables().items():
        xlsx_path = DATA_DIR / f"{sheet}.xlsx"
        layout = schema_layout(sheet)
        if any(f.width != 1 for f in layout):
            fail(f"{sheet} 的 schema 里出现了非标量列,new-tables 只会写全标量表,先同步脚本")
        names = [f.name for f in layout]
        for row in rows:
            if list(row.keys()) != names:
                fail(f"{sheet} 数据行的键 {list(row.keys())} 与 schema 字段 {names} 不一致 —— schema 改过列,先同步本脚本")
        ids = [row["id"] for row in rows]
        if len(set(ids)) != len(ids):
            fail(f"{sheet} 数据行 id 重复")

        wanted = [[None] * len(names) for _ in range(DATA_BEGIN_ROW - 1 + len(rows))]
        wanted[HEADER_ROW - 1] = list(names)
        wanted[COMMENT_ROW - 1] = [f.comment or None for f in layout]
        for offset, row in enumerate(rows):
            wanted[DATA_BEGIN_ROW - 1 + offset] = [row[name] for name in names]

        if xlsx_path.exists():
            ws = only_sheet(openpyxl.load_workbook(xlsx_path), xlsx_path)
            if ws.title == sheet and first_difference(sheet_matrix(ws), wanted) is None:
                print(f"[已是目标状态] {rel(xlsx_path)} 已存在且内容一致({len(rows)} 行),不改文件。")
                continue
            fail(f"{rel(xlsx_path)} 已存在但内容与本脚本不同(可能策划已调过数值)—— 本脚本不覆盖。")

        print(f"将新建 {rel(xlsx_path)}:sheet={sheet!r},{len(names)} 列 × {len(rows)} 行数据")
        print(f"  列:{names}")
        print(f"  首行:{[rows[0][n] for n in names]}  末行:{[rows[-1][n] for n in names]}")
        if dry_run:
            continue
        wb = openpyxl.Workbook()
        ws = wb.active
        # 首个 sheet 的名字必须等于 cfg_sheet:导表器按它匹配 schema,默认的 "Sheet" 会被派到旧表头路径而失败。
        ws.title = sheet
        for r, values in enumerate(wanted, start=1):
            for c, value in enumerate(values, start=1):
                if value is not None:
                    ws.cell(row=r, column=c, value=value)
        wb.save(xlsx_path)

        check_ws = only_sheet(openpyxl.load_workbook(xlsx_path), xlsx_path, sheet)
        diff = first_difference(sheet_matrix(check_ws), wanted)
        if diff:
            fail(f"保存后读回的 {xlsx_path.name} 与写入内容不一致({diff})—— 删除该文件后重试")
        bind(sheet, xlsx_path)
        print(f"[完成] {rel(xlsx_path)}")
    if dry_run:
        print("[dry-run] 未写文件。")


# ════════════════════════════════════════════════════════════════════════════
# ② item-columns
# ════════════════════════════════════════════════════════════════════════════

ITEM_SHEET = "Item"
ITEM_XLSX = DATA_DIR / "Item.xlsx"


def cmd_item_columns(dry_run: bool) -> None:
    layout = {f.name: f for f in schema_layout(ITEM_SHEET)}
    missing_in_schema = [n for n in ITEM_NEW_COLUMNS if n not in layout]
    if missing_in_schema:
        fail(f"{schema_path(ITEM_SHEET).name} 里没有字段 {missing_in_schema} —— schema 还没加这些字段")
    if layout[ITEM_NEW_COLUMNS[-1]].width != 1:
        fail(f"ITEM_NEW_COLUMNS 的最后一项 {ITEM_NEW_COLUMNS[-1]!r} 不是标量:最右一列必须是标量(见常量表上方的说明)")

    wb = openpyxl.load_workbook(ITEM_XLSX)
    ws = only_sheet(wb, ITEM_XLSX, ITEM_SHEET)
    before = sheet_matrix(ws)
    header = header_cells(ws)

    present = [n for n in ITEM_NEW_COLUMNS if n in header]
    if len(present) == len(ITEM_NEW_COLUMNS):
        bind(ITEM_SHEET, ITEM_XLSX)
        print(f"[已是目标状态] {rel(ITEM_XLSX)} 已有 {list(ITEM_NEW_COLUMNS)} 且与 schema 绑定通过,不改文件。")
        return
    if present:
        fail(f"{rel(ITEM_XLSX)} 只有一部分新列 {present}(缺 {[n for n in ITEM_NEW_COLUMNS if n not in header]})"
             f" —— 不是本脚本写出来的状态,先查清再手工补齐。")
    unknown = [n for n in header if n not in layout]
    if unknown:
        fail(f"{rel(ITEM_XLSX)} 第 1 行有 schema 里没有的列 {unknown}")
    if max(header.values()) != ws.max_column:
        fail(f"{rel(ITEM_XLSX)} 的最右一列(第 {ws.max_column} 列)没有表头 —— 现有最后一个字段不是标量,"
             f"或表右侧有游离数据;追加新列会把它们并进那个字段,先修表。")

    patches: list[Patch] = []
    col = ws.max_column + 1
    plan = []
    for name in ITEM_NEW_COLUMNS:
        field = layout[name]
        patches.append((HEADER_ROW, col, name))        # 数组 / 子消息列只在首格写名字,后续格第 1 行留空
        if field.comment:
            patches.append((COMMENT_ROW, col, field.comment))
        last = col + field.width - 1
        span = get_column_letter(col) if field.width == 1 else f"{get_column_letter(col)}..{get_column_letter(last)}"
        plan.append(f"  {span:<6} {name}" + (f"(每条 {len(field.sub_names)} 列:{'/'.join(field.sub_names)})"
                                             if field.sub_names else ""))
        col += field.width

    print(f"{rel(ITEM_XLSX)}:现有 {ws.max_column} 列、数据到第 {ws.max_row} 行;将在最右追加 {col - 1 - ws.max_column} 个物理列"
          f"(既有列与既有行一格不动,既有行的新列留空 = 默认值):")
    print("\n".join(plan))
    if dry_run:
        print("[dry-run] 未写文件。")
        return
    apply_patches_and_verify(ITEM_XLSX, ITEM_SHEET, wb, ws, before, patches, bind_schema=True)
    print(f"[完成] {rel(ITEM_XLSX)} 共 {col - 1} 列,与 schema 绑定通过。")


# ════════════════════════════════════════════════════════════════════════════
# ③ item-rows
# ════════════════════════════════════════════════════════════════════════════

def sample_equip_rows() -> list[dict]:
    rows = []
    for part in SAMPLE_EQUIP_PARTS:
        for tier, level in enumerate(EQUIP_LEVEL_TIERS, start=1):
            icon = WEAPON_ICON_DIR + part["icons"][tier - 1] if part["icons"] else None
            rows.append(dict(
                id=1000 + part["order"] * 100 + tier,
                max_stack_size=1,                 # 装备不可堆叠:可堆叠会把不同属性的实例并成一堆(§2.5)
                equip_kind=part["equip_kind"],
                battle_usable=0,
                battle_heal_hp=0,
                battle_heal_mp=0,
                name=part["names"][tier - 1],
                description=part["description"].format(level=level),
                icon_key=icon,
                equip_level=level,
                equip_class=0,
                affix_pool=part["affix_pool"],
                affix_rule=SAMPLE_AFFIX_RULE,
                base_attr=[(attr_id(name), per_level * level + const)
                           for name, per_level, const in part["base_attr"]],
            ))
    return rows


def item_row_cells(row: dict, layout: list[FieldLayout], header: dict[str, int]) -> dict[int, object]:
    """一行样例装备 -> {列号: 值}。子消息数组用不到的槽不出现在结果里(留空;填 0 会被导成一条 id=0 的基础属性)。"""
    cells: dict[int, object] = {}
    for field in layout:
        col = header[field.name]
        value = row[field.name]
        if not field.sub_names:
            if value is not None:
                cells[col] = value
            continue
        per = len(field.sub_names)
        if len(value) * per > field.width:
            fail(f"物品 {row['id']} 的 {field.name} 有 {len(value)} 条,超过 (cfg_slots) = {field.width // per}")
        for index, entry in enumerate(value):
            if len(entry) != per:
                fail(f"物品 {row['id']} 的 {field.name} 第 {index + 1} 条应有 {per} 个值")
            for offset, sub_value in enumerate(entry):
                cells[col + index * per + offset] = sub_value
    return cells


def cmd_item_rows(dry_run: bool) -> None:
    if not (DATA_DIR / "EquipAffixRule.xlsx").exists():
        fail("data/EquipAffixRule.xlsx 还不存在 —— 样例装备的 affix_rule 是指向它的外键,先执行 new-tables")
    layout = schema_layout(ITEM_SHEET)
    rows = sample_equip_rows()
    for row in rows:
        if set(row.keys()) != {f.name for f in layout}:
            fail(f"样例装备的键与 {schema_path(ITEM_SHEET).name} 的字段不一致 —— schema 改过列,先同步本脚本:"
                 f"多 {sorted(set(row) - {f.name for f in layout})},少 {sorted({f.name for f in layout} - set(row))}")

    wb = openpyxl.load_workbook(ITEM_XLSX)
    ws = only_sheet(wb, ITEM_XLSX, ITEM_SHEET)
    before = sheet_matrix(ws)
    header = header_cells(ws)
    if any(f.name not in header for f in layout):
        fail(f"{rel(ITEM_XLSX)} 还没有列 {[f.name for f in layout if f.name not in header]} —— 先执行 item-columns")
    bind(ITEM_SHEET, ITEM_XLSX)

    id_col = header["id"]
    existing: dict[int, int] = {}       # id -> 行号
    last_data_row = DATA_BEGIN_ROW - 1
    for row_idx in range(DATA_BEGIN_ROW, ws.max_row + 1):
        value = ws.cell(row=row_idx, column=id_col).value
        if value is None or value == "":
            if any(v is not None for v in before[row_idx - 1]):
                fail(f"{rel(ITEM_XLSX)} 第 {row_idx} 行没有 id 却有数据 —— 先修表")
            continue
        if not isinstance(value, (int, float)) or int(value) != value:
            fail(f"{rel(ITEM_XLSX)} 第 {row_idx} 行的 id 不是整数:{value!r}")
        if int(value) in existing:
            fail(f"{rel(ITEM_XLSX)} 里 id {int(value)} 出现了两次(第 {row_idx} 行)—— 先修表")
        existing[int(value)] = row_idx
        last_data_row = row_idx

    wanted_cells = {row["id"]: item_row_cells(row, layout, header) for row in rows}
    taken = [row["id"] for row in rows if row["id"] in existing]
    if taken:
        if len(taken) == len(rows):
            for row in rows:
                row_idx = existing[row["id"]]
                actual = {c: v for c, v in enumerate(before[row_idx - 1], start=1) if v is not None}
                if actual != wanted_cells[row["id"]]:
                    fail(f"{rel(ITEM_XLSX)} 第 {row_idx} 行(id {row['id']})已存在,但内容与本脚本不同"
                         f"(可能策划已调过数值)—— 本脚本不覆盖。")
            print(f"[已是目标状态] {rel(ITEM_XLSX)} 已有这 {len(rows)} 件样例装备且内容一致,不改文件。")
            return
        fail(f"{rel(ITEM_XLSX)} 里这些 id 已被占用:{taken} —— 与样例装备的 id 段冲突(或只写进了一部分),先查清。")

    patches: list[Patch] = []
    print(f"{rel(ITEM_XLSX)}:现有 {len(existing)} 行数据,末行是第 {last_data_row} 行;将追加 {len(rows)} 行"
          f"(id {rows[0]['id']}..{rows[-1]['id']},与现有 id 无冲突):")
    for offset, row in enumerate(rows, start=1):
        row_idx = last_data_row + offset
        patches.extend((row_idx, col, value) for col, value in sorted(wanted_cells[row["id"]].items()))
        print(f"  第 {row_idx} 行: id={row['id']} {row['name']} kind={row['equip_kind']} Lv{row['equip_level']} "
              f"pool={row['affix_pool']} rule={row['affix_rule']} base_attr={row['base_attr']} "
              f"icon={row['icon_key'] or '(空)'}")
    if dry_run:
        print("[dry-run] 未写文件。")
        return
    apply_patches_and_verify(ITEM_XLSX, ITEM_SHEET, wb, ws, before, patches, bind_schema=True)
    print(f"[完成] {rel(ITEM_XLSX)} 数据行 {len(existing)} -> {len(existing) + len(rows)},原有行逐格一致。")


# ════════════════════════════════════════════════════════════════════════════
# ④ equip-slots
# ════════════════════════════════════════════════════════════════════════════

SLOT_SHEET = "EquipSlot"
SLOT_XLSX = DATA_DIR / "EquipSlot.xlsx"
SLOT_COLUMNS = ("id", "equip_kind", "name")


def cmd_equip_slots(dry_run: bool) -> None:
    layout = schema_layout(SLOT_SHEET)
    if tuple(f.name for f in layout) != SLOT_COLUMNS:
        fail(f"{schema_path(SLOT_SHEET).name} 的字段应为 {list(SLOT_COLUMNS)},实际 {[f.name for f in layout]} —— 先同步本脚本")
    name_comment = layout[2].comment

    wb = openpyxl.load_workbook(SLOT_XLSX)
    ws = only_sheet(wb, SLOT_XLSX, SLOT_SHEET)
    before = sheet_matrix(ws)
    header = header_cells(ws)
    if any(n not in SLOT_COLUMNS for n in header) or "id" not in header or "equip_kind" not in header:
        fail(f"{rel(SLOT_XLSX)} 第 1 行应为 id / equip_kind(/ name),实际 {list(header)}")

    patches: list[Patch] = []
    plan: list[str] = []
    if "name" in header:
        name_col = header["name"]
    else:
        if max(header.values()) != ws.max_column:
            fail(f"{rel(SLOT_XLSX)} 的最右一列没有表头,先修表")
        name_col = ws.max_column + 1
        patches.append((HEADER_ROW, name_col, "name"))
        if name_comment:
            patches.append((COMMENT_ROW, name_col, name_comment))
        plan.append(f"  在第 {get_column_letter(name_col)} 列追加表头 name(既有行的 name 留空)")

    existing: dict[int, int] = {}
    last_data_row = DATA_BEGIN_ROW - 1
    for row_idx in range(DATA_BEGIN_ROW, ws.max_row + 1):
        value = ws.cell(row=row_idx, column=header["id"]).value
        if value is None or value == "":
            if any(v is not None for v in before[row_idx - 1]):
                fail(f"{rel(SLOT_XLSX)} 第 {row_idx} 行没有 id 却有数据 —— 先修表")
            continue
        if not isinstance(value, (int, float)) or int(value) != value:
            fail(f"{rel(SLOT_XLSX)} 第 {row_idx} 行的 id 不是整数:{value!r}")
        if int(value) in existing:
            fail(f"{rel(SLOT_XLSX)} 里槽号 {int(value)} 出现了两次 —— 先修表")
        existing[int(value)] = row_idx
        last_data_row = row_idx

    taken = [slot for slot, _k, _n in EQUIP_SLOTS if slot in existing]
    if taken and len(taken) != len(EQUIP_SLOTS):
        fail(f"{rel(SLOT_XLSX)} 里槽号 {taken} 已存在但不是完整的 4 行 —— 先查清是谁加的。")
    if taken:
        if "name" not in header:
            fail(f"{rel(SLOT_XLSX)} 已有槽 {taken} 却没有 name 列 —— 不是本脚本写出来的状态,先查清。")
        for slot, kind, name in EQUIP_SLOTS:
            row_idx = existing[slot]
            actual = (ws.cell(row=row_idx, column=header["equip_kind"]).value, ws.cell(row=row_idx, column=name_col).value)
            if actual != (kind, name):
                fail(f"{rel(SLOT_XLSX)} 第 {row_idx} 行(槽 {slot})是 {actual},与要写的 {(kind, name)} 不同 —— 本脚本不覆盖。")
    else:
        for offset, (slot, kind, name) in enumerate(EQUIP_SLOTS, start=1):
            row_idx = last_data_row + offset
            patches += [(row_idx, header["id"], slot), (row_idx, header["equip_kind"], kind), (row_idx, name_col, name)]
            plan.append(f"  第 {row_idx} 行: 槽 {slot} -> 部位 {kind}({name})")

    if not patches:
        bind(SLOT_SHEET, SLOT_XLSX)
        print(f"[已是目标状态] {rel(SLOT_XLSX)} 已有 name 列与槽 {[s for s, _k, _n in EQUIP_SLOTS]},不改文件。")
        return
    print(f"{rel(SLOT_XLSX)}:现有槽 {sorted(existing)}(一格不动);将:")
    print("\n".join(plan))
    if dry_run:
        print("[dry-run] 未写文件。")
        return
    apply_patches_and_verify(SLOT_XLSX, SLOT_SHEET, wb, ws, before, patches, bind_schema=True)
    print(f"[完成] {rel(SLOT_XLSX)} 槽位 {len(existing)} -> {len(existing) + (0 if taken else len(EQUIP_SLOTS))},原有行逐格一致。")


# ════════════════════════════════════════════════════════════════════════════
# ⑤ tip-codes
# ════════════════════════════════════════════════════════════════════════════

TIP_XLSX = DATA_DIR / "tip" / "Tip.xlsx"
TIP_STATE = EXPORTER_DIR / "state" / "mapping" / "tip_enum_ids" / "tip_enum_ids.json"
TIP_DATA_BEGIN_ROW = 18            # 与导表器 enum_gen._TIP_DATA_BEGIN_ROW 一致
TIP_PREFIX = "//"


def is_group_header(name: str) -> bool:
    """组头 vs 说明行,判据与导表器一致(enum_gen.py):"//" 后**紧贴**非空白字符 = 组头,
    "// 这是一句说明" = 说明行。不能用"有没有 base="来分,那会与导表器的口径分叉。"""
    return name.startswith(TIP_PREFIX) and len(name) > 2 and not name[2].isspace()


def parse_group_header(name: str, row_idx: int) -> tuple[str, int, int]:
    parts = name[2:].split()
    kv = dict(tok.split("=", 1) for tok in parts[1:] if "=" in tok)
    try:
        return parts[0], int(kv["base"]), int(kv["width"])
    except (KeyError, ValueError):
        fail(f"{TIP_XLSX.name} 第 {row_idx} 行 {name!r} 像组头却没有合法的 base= / width= —— 导表器也会拒绝,先修表")


def tip_parse(matrix: list[list], fault_col: int):
    """返回 (组列表, 说明行列表)。组 = dict(name, base, width, row, codes=[(码名, 文案, fault 原值)])。"""
    groups: list[dict] = []
    notes: list[tuple[int, str]] = []
    for row_idx in range(TIP_DATA_BEGIN_ROW, len(matrix) + 1):
        row = matrix[row_idx - 1]
        a = row[0]
        if not isinstance(a, str) or not a.strip():
            continue
        name = a.strip()
        if name.startswith(TIP_PREFIX):
            if not is_group_header(name):
                notes.append((row_idx, name))
                continue
            group, base, width = parse_group_header(name, row_idx)
            groups.append(dict(name=group, base=base, width=width, row=row_idx, codes=[]))
            continue
        if not groups:
            fail(f"{TIP_XLSX.name} 第 {row_idx} 行的码 {name!r} 出现在第一个组头之前")
        fault = row[fault_col - 1] if fault_col <= len(row) else None
        groups[-1]["codes"].append((name, row[1] if len(row) > 1 else None, fault))
    return groups, notes


def cmd_tip_codes(dry_run: bool) -> None:
    wb = openpyxl.load_workbook(TIP_XLSX)
    ws = only_sheet(wb, TIP_XLSX)
    before = sheet_matrix(ws)
    header = [str(v).strip().lower() if v is not None else "" for v in before[HEADER_ROW - 1]]
    if header[:2] != ["name", "commen"] or header.count("fault") != 1:
        fail(f"{TIP_XLSX.name} 第 1 行应为 name / commen / … / fault(fault 恰好一列),实际 {before[HEADER_ROW - 1]}")
    fault_col = header.index("fault") + 1

    groups, notes = tip_parse(before, fault_col)
    header_text = f"//{TIP_EQUIP_GROUP} base={TIP_EQUIP_BASE} width={TIP_EQUIP_WIDTH}"
    wanted = [(name, text, 1 if fault else None) for name, text, fault in TIP_EQUIP_CODES]
    new_names = [name for name, _t, _f in TIP_EQUIP_CODES]

    mine = next((g for g in groups if g["name"] == TIP_EQUIP_GROUP), None)
    if mine is not None:
        # 已是目标状态 = 段声明一致,且这 11 个码是本组开头的连续一段。不要求恰好在组尾:以后还会往组尾追加。
        if (mine["base"], mine["width"]) == (TIP_EQUIP_BASE, TIP_EQUIP_WIDTH) and mine["codes"][:len(wanted)] == wanted:
            print(f"[已是目标状态] {header_text}(第 {mine['row']} 行)下已有这 {len(wanted)} 个码(码名 / 文案 / fault 一致),不改文件。")
            return
        fail(f"{TIP_XLSX.name} 第 {mine['row']} 行已有 //{TIP_EQUIP_GROUP} 组,但段声明或开头 {len(wanted)} 个码与本脚本不同 —— 先查清是谁加的。")

    # 以下三项都是「先核对再动手」:码名全表唯一、28000 段没被占、也没被说明行预留。
    owner = {code: g["name"] for g in groups for code, _t, _f in g["codes"]}
    clash = [f"{n}(在 {owner[n]})" for n in new_names if n in owner]
    if clash:
        fail(f"这些码名已在表里:{clash} —— 码名全局唯一(生成的 tip proto 没有 package),先查清。")
    lo, hi = TIP_EQUIP_BASE, TIP_EQUIP_BASE + TIP_EQUIP_WIDTH
    overlap = [f"{g['name']} [{g['base']},{g['base'] + g['width']})" for g in groups
               if g["base"] < hi and lo < g["base"] + g["width"]]
    if overlap:
        fail(f"段 [{lo},{hi}) 已被占用:{overlap} —— 改用下一个空闲整千段(改 TIP_EQUIP_BASE),"
             f"并同步 docs/design/equipment-attributes.md §2.7 与 §10。")
    reserved = [f"第 {row} 行 {text!r}" for row, text in notes if str(TIP_EQUIP_BASE) in text]
    if reserved:
        fail(f"{TIP_EQUIP_BASE} 段已被说明行预留:{reserved} —— 改用下一个空闲整千段并同步设计文档 §2.7 与 §10。")
    if TIP_STATE.exists():
        # 发号 state 里的整组是墓碑:表里删掉的组仍占着段。只读核对,不写 state。
        state_groups = json.loads(TIP_STATE.read_text(encoding="utf-8")).get("groups", {})
        tomb = [f"{n} [{g['base']},{g['base'] + g['width']})" for n, g in state_groups.items()
                if int(g["base"]) < hi and lo < int(g["base"]) + int(g["width"])]
        if tomb:
            fail(f"发号 state({rel(TIP_STATE)})里段 [{lo},{hi}) 已属于 {tomb} —— 换一个段,不要手改 state。")

    last_row = max((r for r, row in enumerate(before, start=1) if any(v is not None for v in row)), default=0)
    header_row = last_row + 2          # 与既有组一样,组头前留一个空行
    patches: list[Patch] = [(header_row, 1, header_text)]
    print(f"{TIP_XLSX.name}:现有 {len(groups)} 个组、{len(owner)} 个码,末行是第 {last_row} 行;"
          f"段 [{lo},{hi}) 空闲、{len(new_names)} 个码名全表唯一。将追加:")
    print(f"  第 {header_row} 行: {header_text}")
    for offset, (name, text, fault) in enumerate(wanted, start=1):
        row_idx = header_row + offset
        patches += [(row_idx, 1, name), (row_idx, 2, text)]
        if fault is not None:
            patches.append((row_idx, fault_col, fault))
        print(f"  第 {row_idx} 行: {name} | {text}" + (" | fault=1" if fault else ""))
    if dry_run:
        print("[dry-run] 未写文件。")
        return
    apply_patches_and_verify(TIP_XLSX, None, wb, ws, before, patches, bind_schema=False)

    check_ws = only_sheet(openpyxl.load_workbook(TIP_XLSX), TIP_XLSX)
    groups_after, _notes = tip_parse(sheet_matrix(check_ws), fault_col)
    if [(g["name"], g["codes"]) for g in groups_after[:-1]] != [(g["name"], g["codes"]) for g in groups] \
            or groups_after[-1]["name"] != TIP_EQUIP_GROUP or groups_after[-1]["codes"] != wanted:
        fail(f"保存后各组的码与预期不符 —— 立刻 git checkout -- {rel(TIP_XLSX)} 还原")
    print(f"[完成] 组 {len(groups)} -> {len(groups_after)},原有 {len(owner)} 个码逐格一致。"
          f"导表后核对:这 {len(wanted)} 个码的 id 应为 {lo}..{lo + len(wanted) - 1},tip_enum_ids.json 只增不改;"
          f"新域的 proto/tip/{TIP_EQUIP_GROUP}_tip.pb.{{h,cc}} 要登记进 cpp/generated/table 的 CMakeLists.txt 与 table.vcxproj。")


# ════════════════════════════════════════════════════════════════════════════
# ⑥ message-limiter
# ════════════════════════════════════════════════════════════════════════════

MESSAGE_ID_TXT = REPO_ROOT / "proto" / "message_id.txt"
LIMITER_SHEET = "MessageLimiter"
LIMITER_XLSX = DATA_DIR / "MessageLimiter.xlsx"
LIMITER_COLUMNS = ("id", "max_requests", "time_window", "tip_message")


def resolve_message_ids() -> dict[str, int]:
    """方法名 -> 消息号。任何一个方法查不到就中止 —— 此时还没有打开过 xlsx,什么都不会写。"""
    if not MESSAGE_ID_TXT.is_file():
        fail(f"找不到 {rel(MESSAGE_ID_TXT)} —— 消息号表由全量 proto-gen 生成,先跑 proto-gen。未写任何文件。")
    ids: dict[str, int] = {}
    for line_no, line in enumerate(MESSAGE_ID_TXT.read_text(encoding="utf-8").splitlines(), start=1):
        line = line.strip()
        if not line:
            continue
        number, sep, key = line.partition("=")
        if not sep or not number.isdigit():
            fail(f"{MESSAGE_ID_TXT.name}:{line_no} 不是 `<号>=<服务名+方法名>` 形状:{line!r}")
        ids[key] = int(number)
    methods = list(LIMITER_NEW_METHODS) + list(LIMITER_BACKFILL_METHODS)
    missing = [LIMITER_SERVICE + m for m in methods if LIMITER_SERVICE + m not in ids]
    if missing:
        fail(f"{MESSAGE_ID_TXT.name} 里找不到 {missing} —— 全量 proto-gen 还没跑(或没跑成功),消息号未定,"
             f"现在不能填档位。未写任何文件。")
    return {m: ids[LIMITER_SERVICE + m] for m in methods}


def cmd_message_limiter(dry_run: bool) -> None:
    message_ids = resolve_message_ids()

    wb = openpyxl.load_workbook(LIMITER_XLSX)
    ws = only_sheet(wb, LIMITER_XLSX, LIMITER_SHEET)
    before = sheet_matrix(ws)
    header = header_cells(ws)
    for name in LIMITER_COLUMNS:
        if name not in header:
            fail(f"{LIMITER_XLSX.name} 第 1 行找不到列 {name!r},实际表头 {list(header)}")

    existing: dict[int, tuple] = {}
    last_data_row = DATA_BEGIN_ROW - 1
    for row_idx in range(DATA_BEGIN_ROW, ws.max_row + 1):
        value = ws.cell(row=row_idx, column=header["id"]).value
        if value is None or value == "":
            continue
        if not isinstance(value, (int, float)) or int(value) != value:
            fail(f"{LIMITER_XLSX.name} 第 {row_idx} 行的 id 不是整数:{value!r}")
        if int(value) in existing:
            fail(f"{LIMITER_XLSX.name} 里消息号 {int(value)} 出现了两次(第 {row_idx} 行)—— 先修表")
        existing[int(value)] = tuple(ws.cell(row=row_idx, column=header[name]).value for name in LIMITER_COLUMNS[1:])
        last_data_row = row_idx

    to_append = []
    for method, max_requests in {**LIMITER_NEW_METHODS, **LIMITER_BACKFILL_METHODS}.items():
        message_id = message_ids[method]
        wanted = (max_requests, LIMITER_TIME_WINDOW, LIMITER_TIP_MESSAGE)
        if message_id not in existing:
            to_append.append((message_id, method, wanted))
        elif method in LIMITER_BACKFILL_METHODS:
            if existing[message_id] != wanted:
                print(f"  [保留] 消息号 {message_id}({LIMITER_SERVICE}{method})已在表里,档位 {existing[message_id]},不改。")
        elif existing[message_id] != wanted:
            # 最可能的原因:这个号以前属于别的方法(被释放的旧号易主),表里留着旧方法的档位。
            fail(f"消息号 {message_id}({LIMITER_SERVICE}{method})在表里已有一行 {existing[message_id]},"
                 f"与要写的 {wanted} 不同 —— 先确认那一行是谁的,本脚本不覆盖。")

    if not to_append:
        print(f"[已是目标状态] {len(message_ids)} 个背包 / 装备消息号都已在表里,不改文件。")
        return

    patches: list[Patch] = []
    print(f"现有数据行 {len(existing)} 条,末行是第 {last_data_row} 行;将追加 {len(to_append)} 行:")
    for offset, (message_id, method, wanted) in enumerate(to_append, start=1):
        row_idx = last_data_row + offset
        patches.append((row_idx, header["id"], message_id))
        patches += [(row_idx, header[name], value) for name, value in zip(LIMITER_COLUMNS[1:], wanted)]
        print(f"  第 {row_idx} 行: id={message_id:<4} {LIMITER_SERVICE}{method:<12} "
              f"max_requests={wanted[0]} time_window={wanted[1]} tip_message={wanted[2]}")
    if dry_run:
        print("[dry-run] 未写文件。")
        return
    apply_patches_and_verify(LIMITER_XLSX, LIMITER_SHEET, wb, ws, before, patches, bind_schema=True)
    print(f"[完成] 数据行 {len(existing)} -> {len(existing) + len(to_append)},原有行逐格一致。接着再跑一次导表器;"
          f"通过标准:generated/tables/messagelimiter.json 里能查到上面这些 id。部署时 gate 要重载 MessageLimiter。")


COMMANDS = {
    "new-tables": cmd_new_tables,
    "item-columns": cmd_item_columns,
    "item-rows": cmd_item_rows,
    "equip-slots": cmd_equip_slots,
    "tip-codes": cmd_tip_codes,
    "message-limiter": cmd_message_limiter,
}


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("command", choices=tuple(COMMANDS))
    parser.add_argument("--dry-run", action="store_true", help="只打印将要做什么,不写文件")
    args = parser.parse_args()
    self_check()
    COMMANDS[args.command](args.dry_run)


if __name__ == "__main__":
    # Windows 默认控制台编码(cp936 / cp1252)打不出部分字符时会在 print 上抛 UnicodeEncodeError,
    # 而那时文件可能已经写完 —— 统一成 utf-8,免得"改成功了却以非 0 退出"误导人。
    for stream in (sys.stdout, sys.stderr):
        if hasattr(stream, "reconfigure"):
            stream.reconfigure(encoding="utf-8", errors="replace")
    main()
