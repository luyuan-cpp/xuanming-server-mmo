#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""帮会二期 B5a 要对 xlsx 做的三件事,写成可重复执行的脚本。

为什么是脚本而不是手改:
  - data/tip/Tip.xlsx、data/MessageLimiter.xlsx 同时有多个会话在写,它们是二进制、不能 3-way 合并,
    整表写回会**静默**删掉别人的行(导表器照常跑通,只是少发几个号)。唯一安全的写法是 openpyxl
    从磁盘重新载入 → 只在自己那一段 insert_rows / 追加 → save → 读回来逐段核对。
  - MessageLimiter 按**消息号**登记档位,而新方法落到哪个号要等全量 proto-gen 之后才定,
    所以这里按方法名去 proto/message_id.txt 查号,不把号写死。
  - 两张新配表的 schema(data/schema/guilddonate_table.proto、guildshop_table.proto)已落盘。
    **导表器要求 schema 与 xlsx 成对**:只有 schema 没有 xlsx 时,任何人跑 `dev.bat gen` 都会以
    "这些权威 schema 没有对应的源表" 整批失败。所以 new-tables 必须先于下一次导表执行;
    生成的两个 xlsx 只在一台机器上生成并提交(openpyxl 每次写出的字节都不同,两台机器各生成一份
    会在 git 里撞成二进制 add/add 冲突),其它机器等 pull。

三个子命令都**幂等**:已经是目标状态就只报告、不改文件;发现与预期不符的现状就中止,绝不覆盖。

用法(仓库根目录执行;需要 Python 3 + openpyxl,`pip install openpyxl`):

    python tools/scripts/guild_b5a_xlsx_patch.py new-tables        # ① 任何导表之前
    python tools/scripts/guild_b5a_xlsx_patch.py tip-codes         # ② 导表之前(与其它会话错开,见下)
    python tools/scripts/guild_b5a_xlsx_patch.py message-limiter   # ③ 全量 proto-gen 之后、第二次导表之前

加 --dry-run 只打印将要做什么,不写文件。tip-codes / message-limiter 执行前先确认
`git status --short data/tip/Tip.xlsx data/MessageLimiter.xlsx` 没有别人的未提交改动(90 清单 G-08:
任一时刻只允许一方有未提交修改)。

默认数值出处:docs/design/guild-phase2/README.md §4、05-economy.md §5.10 / §5.12 / §5.13,
并按 90-consistency.md 的 X-04(GuildShop 204 的等级要求是 6)、Y-05(tip 顺序)订正。
"""

from __future__ import annotations

import argparse
import re
import sys
from pathlib import Path

try:
    import openpyxl
except ImportError:  # pragma: no cover - 环境问题,给一句能照着做的话
    sys.exit("缺少 openpyxl:请先执行  pip install openpyxl")

REPO_ROOT = Path(__file__).resolve().parents[2]
DATA_DIR = REPO_ROOT / "data"
SCHEMA_DIR = DATA_DIR / "schema"

# 主表版式(data/AGENTS.md):第 1 行列名 / 第 2-4 行留空 / 第 5 行中文说明(schema 注释首行的投影)/ 第 6 行起数据。
HEADER_ROW = 1
COMMENT_ROW = 5
DATA_BEGIN_ROW = 6


def fail(message: str) -> "NoReturn":
    sys.exit(f"[中止] {message}")


def only_sheet(wb, path: Path):
    if len(wb.sheetnames) != 1:
        fail(f"{path.name} 预期只有一个 sheet,实际 {wb.sheetnames}")
    return wb[wb.sheetnames[0]]


# ── new-tables ───────────────────────────────────────────────────────────────

# 列名与列序不在这里写第二份:从权威 schema 读。这里只给数据,键必须与 schema 字段逐一对上。
NEW_TABLES = {
    "GuildDonate": {
        "schema": "guilddonate_table.proto",
        "rows": [
            dict(id=1, name="银两小捐", currency_type=0, cost_amount=10000, contribution_gain=10,
                 funds_gain=1000, daily_limit=5, min_guild_level=1),
            dict(id=2, name="银两大捐", currency_type=0, cost_amount=100000, contribution_gain=120,
                 funds_gain=12000, daily_limit=2, min_guild_level=1),
            dict(id=3, name="灵石捐献", currency_type=1, cost_amount=100, contribution_gain=200,
                 funds_gain=20000, daily_limit=1, min_guild_level=1),
        ],
    },
    "GuildShop": {
        "schema": "guildshop_table.proto",
        # 物品 id 均在 Item 表内:15-23 的 max_stack_size 为 999,12、13 为 1(所以 202 / 204 每份只能是 1 个)。
        "rows": [
            dict(id=101, name="培元丹", category=1, item_id=15, item_count=5, cost_contribution=30,
                 required_guild_level=1, limit_period=1, limit_count=10),
            dict(id=102, name="回灵散", category=1, item_id=16, item_count=5, cost_contribution=30,
                 required_guild_level=1, limit_period=1, limit_count=10),
            dict(id=103, name="精炼石", category=1, item_id=17, item_count=1, cost_contribution=80,
                 required_guild_level=2, limit_period=1, limit_count=5),
            dict(id=104, name="修行秘录残页", category=1, item_id=18, item_count=1, cost_contribution=150,
                 required_guild_level=3, limit_period=2, limit_count=5),
            dict(id=201, name="帮会令牌", category=2, item_id=19, item_count=1, cost_contribution=300,
                 required_guild_level=3, limit_period=2, limit_count=3),
            dict(id=202, name="玄铁护符", category=2, item_id=12, item_count=1, cost_contribution=800,
                 required_guild_level=4, limit_period=2, limit_count=1),
            dict(id=203, name="灵兽口粮", category=2, item_id=20, item_count=10, cost_contribution=120,
                 required_guild_level=2, limit_period=1, limit_count=3),
            dict(id=204, name="藏经阁手札", category=2, item_id=13, item_count=1, cost_contribution=1500,
                 required_guild_level=6, limit_period=2, limit_count=1),
            dict(id=301, name="花灯", category=3, item_id=21, item_count=1, cost_contribution=50,
                 required_guild_level=1, limit_period=1, limit_count=5),
            dict(id=302, name="月饼礼盒", category=3, item_id=22, item_count=1, cost_contribution=100,
                 required_guild_level=1, limit_period=2, limit_count=7),
            dict(id=303, name="同心结", category=3, item_id=23, item_count=1, cost_contribution=200,
                 required_guild_level=5, limit_period=0, limit_count=0),
        ],
    },
}

SCHEMA_SHEET_RE = re.compile(r'option\s*\(cfg_sheet\)\s*=\s*"([^"]+)"')
SCHEMA_FIELD_RE = re.compile(r"^\s*(?:repeated\s+)?[\w.]+\s+(\w+)\s*=\s*\d+\s*(?:\[.*\])?\s*;")


def read_schema_columns(schema_path: Path, expect_sheet: str) -> list[tuple[str, str]]:
    """返回 [(字段名, 字段上方注释的首行)],顺序 = schema 声明顺序。只支持本批这种全标量、单行字段的 schema。"""
    text = schema_path.read_text(encoding="utf-8")
    sheet = SCHEMA_SHEET_RE.search(text)
    if not sheet or sheet.group(1) != expect_sheet:
        fail(f"{schema_path.name} 的 cfg_sheet 应为 {expect_sheet!r},实际 {sheet.group(1) if sheet else None!r}")
    columns: list[tuple[str, str]] = []
    pending_comment: list[str] = []
    in_message = False
    for line in text.splitlines():
        stripped = line.strip()
        if stripped.startswith("message "):
            in_message = True
            pending_comment = []
            continue
        if not in_message:
            continue
        if stripped.startswith("//"):
            pending_comment.append(stripped[2:].strip())
            continue
        field = SCHEMA_FIELD_RE.match(line)
        if field:
            columns.append((field.group(1), pending_comment[0] if pending_comment else ""))
        pending_comment = []
    if not columns or columns[0][0] != "id":
        fail(f"{schema_path.name} 解析出的首列应为 id,实际 {columns[:1]}")
    return columns


def sheet_matrix(ws, width: int) -> list[list]:
    return [[ws.cell(row=r, column=c).value for c in range(1, width + 1)] for r in range(1, ws.max_row + 1)]


def cmd_new_tables(dry_run: bool) -> None:
    for sheet_name, spec in NEW_TABLES.items():
        xlsx_path = DATA_DIR / f"{sheet_name}.xlsx"
        columns = read_schema_columns(SCHEMA_DIR / spec["schema"], sheet_name)
        names = [name for name, _ in columns]
        for row in spec["rows"]:
            if list(row.keys()) != names:
                fail(f"{sheet_name} 数据行的键 {list(row.keys())} 与 schema 字段 {names} 不一致 —— schema 改过列,先同步本脚本")
        ids = [row["id"] for row in spec["rows"]]
        if len(set(ids)) != len(ids):
            fail(f"{sheet_name} 数据行 id 重复:{ids}")

        wanted = [[None] * len(names) for _ in range(DATA_BEGIN_ROW - 1 + len(spec["rows"]))]
        wanted[HEADER_ROW - 1] = list(names)
        wanted[COMMENT_ROW - 1] = [comment or None for _, comment in columns]
        for offset, row in enumerate(spec["rows"]):
            wanted[DATA_BEGIN_ROW - 1 + offset] = [row[name] for name in names]

        if xlsx_path.exists():
            ws = only_sheet(openpyxl.load_workbook(xlsx_path), xlsx_path)
            current = sheet_matrix(ws, max(ws.max_column, len(names)))
            padded = [r + [None] * (len(current[0]) - len(r)) for r in wanted] if current else wanted
            if ws.title == sheet_name and current == padded:
                print(f"[已是目标状态] {xlsx_path.relative_to(REPO_ROOT)} 已存在且内容一致,不改文件。")
                continue
            fail(f"{xlsx_path.relative_to(REPO_ROOT)} 已存在但内容与本脚本不同(可能策划已调过数值)—— 本脚本不覆盖。")

        print(f"将新建 {xlsx_path.relative_to(REPO_ROOT)}:sheet={sheet_name!r},{len(names)} 列 × {len(spec['rows'])} 行数据")
        print(f"  列:{names}")
        if dry_run:
            continue
        wb = openpyxl.Workbook()
        ws = wb.active
        # 首个 sheet 的名字必须等于 cfg_sheet:导表器按它匹配 schema,默认的 "Sheet" 会被派到旧表头路径而失败。
        ws.title = sheet_name
        for r, values in enumerate(wanted, start=1):
            for c, value in enumerate(values, start=1):
                if value is not None:
                    ws.cell(row=r, column=c, value=value)
        wb.save(xlsx_path)

        check_ws = only_sheet(openpyxl.load_workbook(xlsx_path), xlsx_path)
        if check_ws.title != sheet_name or sheet_matrix(check_ws, len(names)) != wanted:
            fail(f"保存后读回的 {xlsx_path.name} 与写入内容不一致 —— 删除该文件后重试")
        print(f"[完成] {xlsx_path.relative_to(REPO_ROOT)}")
    if dry_run:
        print("[dry-run] 未写文件。")


# ── tip-codes ────────────────────────────────────────────────────────────────

TIP_XLSX = DATA_DIR / "tip" / "Tip.xlsx"
TIP_GROUP_HEADER_PREFIX = "//"
TIP_GUILD_GROUP = "//guild_error"
# 号由导表器持久化在 tools/data_table_exporter/state/mapping/tip_enum_ids/tip_enum_ids.json:已发的号按码名原样保留,
# 新码名按表内出现顺序取本段最小空位(enum_gen._assign_tip_ids)。所以行序不决定旧码的号;
# 追加在组尾是 90-consistency.md Y-05 的约定(各批按顺序追加),本批 10 个码预期依次拿到 14022..14031。
# 多会话各自导表时真正的冲突点是那份 state 文件:两边给新码发了同一个空位时,后合入的一方重导,不要手改 state。
# 顺序 = Y-05 的 B5a 一行;fault 列全部留空(都是业务拒绝,不是服务端故障)。
# 标点随表内既有行用全角。
TIP_NEW_CODES = [
    ("GuildFundsInsufficient", "帮会资金不足"),
    ("GuildMaxLevel", "帮会已达最高等级"),
    ("GuildDonateLimit", "今日该项捐献次数已用完"),
    ("GuildCurrencyInsufficient", "货币不足，无法捐献"),
    ("GuildAssetPending", "有未结算的帮会资产操作，请稍后再试"),
    ("GuildAssetRejected", "资产结算失败，本次操作已撤销"),
    ("GuildShopGoodsNotFound", "兑换商品不存在"),
    ("GuildShopLevelTooLow", "帮会等级不足"),
    ("GuildShopLimit", "已达限购数量"),
    ("GuildContributionInsufficient", "可用帮贡不足"),
]


def is_group_header(name: str) -> bool:
    """组头 vs 说明行,判据与导表器一致(enum_gen.py):"//" 后**紧贴**非空白字符 = 组头,
    "// 这是一句说明" = 说明行。不能用"有没有 base="来分:那会把漏写 base= 的组头当成说明行,
    它下面的码被静默归进上一组,而导表器那边会报错 —— 两边口径就分叉了。"""
    return name.startswith(TIP_GROUP_HEADER_PREFIX) and len(name) > 2 and not name[2].isspace()


def tip_rows(ws) -> list[tuple[int, str, object, object]]:
    """A 列非空的全部行 [(行号, A, B, C)],含组头与说明行。A 列去首尾空白,B / C 原样。"""
    out = []
    for row_idx, (a, b, c) in enumerate(ws.iter_rows(min_row=1, max_col=3, values_only=True), start=1):
        if isinstance(a, str) and a.strip():
            out.append((row_idx, a.strip(), b, c))
    return out


def tip_groups(rows: list[tuple[int, str, object, object]]) -> dict[str, list[str]]:
    groups: dict[str, list[str]] = {}
    current = None
    for row_idx, name, _b, _c in rows:
        if name.startswith(TIP_GROUP_HEADER_PREFIX):
            if not is_group_header(name):
                continue  # 说明行
            if "base=" not in name:
                fail(f"{TIP_XLSX.name} 第 {row_idx} 行 {name!r} 像组头却没有 base= —— 导表器也会拒绝,先修表")
            current = name.split()[0]
            groups.setdefault(current, [])
            continue
        if current is not None:
            groups[current].append(name)
    return groups


def cmd_tip_codes(dry_run: bool) -> None:
    wb = openpyxl.load_workbook(TIP_XLSX)
    ws = only_sheet(wb, TIP_XLSX)
    rows_before = tip_rows(ws)
    groups_before = tip_groups(rows_before)
    if TIP_GUILD_GROUP not in groups_before:
        fail(f"{TIP_XLSX.name} 里找不到组头 {TIP_GUILD_GROUP}")

    all_codes = [n for codes in groups_before.values() for n in codes]
    new_names = [name for name, _ in TIP_NEW_CODES]
    present = [n for n in new_names if n in all_codes]
    if present:
        # 已是目标状态 = 这 10 个码在本组内连续、顺序一致。不要求它们恰好在组尾:B6a 之后还会往组尾追加活动码。
        guild_codes = groups_before[TIP_GUILD_GROUP]
        if new_names[0] in guild_codes:
            start = guild_codes.index(new_names[0])
            if guild_codes[start:start + len(new_names)] == new_names:
                print(f"[已是目标状态] {TIP_GUILD_GROUP} 组内已有这 {len(new_names)} 个码(连续、顺序一致),不改文件。")
                return
        fail(f"这些码已在表里,但不是连续、按序位于 {TIP_GUILD_GROUP} 组内:{present} —— 码名全局唯一,先查清是谁加的。")

    # 组尾 = 组头之后、下一个 "//" 行(或空行)之前的最后一个码所在行。
    header_row = next(r for r, n, _b, _c in rows_before if is_group_header(n) and n.split()[0] == TIP_GUILD_GROUP)
    last_code_row = header_row
    for row_idx in range(header_row + 1, ws.max_row + 2):
        value = ws.cell(row=row_idx, column=1).value
        if not isinstance(value, str) or not value.strip() or value.strip().startswith(TIP_GROUP_HEADER_PREFIX):
            break
        last_code_row = row_idx
    insert_at = last_code_row + 1
    tail_name = ws.cell(row=last_code_row, column=1).value

    print(f"{TIP_GUILD_GROUP}:现有 {len(groups_before[TIP_GUILD_GROUP])} 个码,组尾是第 {last_code_row} 行 {tail_name!r};"
          f"将在第 {insert_at} 行前插入 {len(TIP_NEW_CODES)} 行:")
    for offset, (name, text) in enumerate(TIP_NEW_CODES):
        print(f"  第 {insert_at + offset} 行: {name} | {text}")
    if dry_run:
        print("[dry-run] 未写文件。")
        return

    ws.insert_rows(insert_at, amount=len(TIP_NEW_CODES))
    for offset, (name, text) in enumerate(TIP_NEW_CODES):
        ws.cell(row=insert_at + offset, column=1, value=name)
        ws.cell(row=insert_at + offset, column=2, value=text)
    wb.save(TIP_XLSX)

    # 读回来逐行核对 A / B / C 三列:原有行一格都不许变(码名、文案、fault 标记仍在同一行上),
    # 只在 insert_at 处多出本批这 10 行。只比码名是不够的 —— fault 标记若与码名错了行,
    # 导表器只在它落到组头 / 空行时才报错,落到相邻的码上就是静默改了故障分类。
    check_ws = only_sheet(openpyxl.load_workbook(TIP_XLSX), TIP_XLSX)
    rows_after = [(a, b, c) for _row, a, b, c in tip_rows(check_ws)]
    expected = ([(a, b, c) for row, a, b, c in rows_before if row < insert_at]
                + [(name, text, None) for name, text in TIP_NEW_CODES]
                + [(a, b, c) for row, a, b, c in rows_before if row >= insert_at])
    if rows_after != expected:
        first_diff = next((i for i, (x, y) in enumerate(zip(rows_after, expected)) if x != y), min(len(rows_after), len(expected)))
        fail(f"保存后读回的内容与预期不符(第一处差异在第 {first_diff + 1} 个非空行;读回 {len(rows_after)} 行,预期 {len(expected)} 行)"
             f" —— 立刻 git checkout -- {TIP_XLSX.relative_to(REPO_ROOT).as_posix()} 还原")
    groups_after = tip_groups(tip_rows(check_ws))
    print(f"[完成] {TIP_GUILD_GROUP} 段 {len(groups_before[TIP_GUILD_GROUP])} -> {len(groups_after[TIP_GUILD_GROUP])} 个码,"
          f"其余 {len(groups_after) - 1} 个段与原有 {len(rows_before)} 行逐格一致。"
          f"导表后核对:这 10 个码的 id 应为 14022..14031,tip_enum_ids.json 只增不改。")


# ── message-limiter ──────────────────────────────────────────────────────────

MESSAGE_ID_TXT = REPO_ROOT / "proto" / "message_id.txt"
LIMITER_XLSX = DATA_DIR / "MessageLimiter.xlsx"
LIMITER_COLUMNS = ("id", "max_requests", "time_window", "tip_message")
GUILD_SERVICE = "GuildService"
# 读 10 次/秒、写 5 次/秒,与表内既有帮会行同档(90-consistency.md G-01);
# time_window / tip_message 照抄既有行(1 秒 / 1000 = kRateLimitExceeded)。
LIMITER_TIERS = {
    "GetGuildDonateOptions": 10,
    "GetGuildShop": 10,
    "DonateToGuild": 5,
    "UpgradeGuild": 5,
    "BuyGuildShopGoods": 5,
}
LIMITER_TIME_WINDOW = 1
LIMITER_TIP_MESSAGE = 1000


def resolve_guild_ids() -> dict[str, int]:
    ids: dict[str, int] = {}
    for line_no, line in enumerate(MESSAGE_ID_TXT.read_text(encoding="utf-8").splitlines(), start=1):
        line = line.strip()
        if not line:
            continue
        number, sep, key = line.partition("=")
        if not sep or not number.isdigit():
            fail(f"{MESSAGE_ID_TXT.name}:{line_no} 不是 `<号>=<服务名+方法名>` 形状:{line!r}")
        ids[key] = int(number)
    missing = [GUILD_SERVICE + m for m in LIMITER_TIERS if GUILD_SERVICE + m not in ids]
    if missing:
        fail(f"{MESSAGE_ID_TXT.name} 里找不到 {missing} —— 全量 proto-gen 还没跑(或没跑成功),消息号未定,现在不能填档位。")
    return {m: ids[GUILD_SERVICE + m] for m in LIMITER_TIERS}


def cmd_message_limiter(dry_run: bool) -> None:
    guild_ids = resolve_guild_ids()

    wb = openpyxl.load_workbook(LIMITER_XLSX)
    ws = only_sheet(wb, LIMITER_XLSX)
    header = [c.value.strip() if isinstance(c.value, str) else c.value for c in ws[HEADER_ROW]]
    col = {}
    for name in LIMITER_COLUMNS:
        if name not in header:
            fail(f"{LIMITER_XLSX.name} 第 1 行找不到列 {name!r},实际表头 {header}")
        col[name] = header.index(name) + 1

    existing: dict[int, tuple] = {}
    last_data_row = DATA_BEGIN_ROW - 1
    for row_idx in range(DATA_BEGIN_ROW, ws.max_row + 1):
        value = ws.cell(row=row_idx, column=col["id"]).value
        if value is None or value == "":
            continue
        if not isinstance(value, (int, float)) or int(value) != value:
            fail(f"{LIMITER_XLSX.name} 第 {row_idx} 行的 id 不是整数:{value!r}")
        message_id = int(value)
        if message_id in existing:
            fail(f"{LIMITER_XLSX.name} 里消息号 {message_id} 出现了两次(第 {row_idx} 行)—— 先修表")
        existing[message_id] = tuple(ws.cell(row=row_idx, column=col[name]).value for name in LIMITER_COLUMNS[1:])
        last_data_row = row_idx

    to_append = []
    for method, max_requests in LIMITER_TIERS.items():
        message_id = guild_ids[method]
        wanted = (max_requests, LIMITER_TIME_WINDOW, LIMITER_TIP_MESSAGE)
        if message_id in existing:
            if existing[message_id] != wanted:
                # 最可能的原因:这个号以前属于别的方法(被释放的旧号易主),表里留着旧方法的档位。
                fail(f"消息号 {message_id}({GUILD_SERVICE}{method})在表里已有一行 {existing[message_id]},"
                     f"与要写的 {wanted} 不同 —— 先确认那一行是谁的,本脚本不覆盖。")
            continue
        to_append.append((message_id, method, wanted))

    if not to_append:
        print(f"[已是目标状态] {len(LIMITER_TIERS)} 个帮会经济消息号都已在表里,不改文件。")
        return

    print(f"现有数据行 {len(existing)} 条,末行是第 {last_data_row} 行;将追加 {len(to_append)} 行:")
    for offset, (message_id, method, wanted) in enumerate(to_append, start=1):
        print(f"  第 {last_data_row + offset} 行: id={message_id:<4} {GUILD_SERVICE}{method:<24} "
              f"max_requests={wanted[0]} time_window={wanted[1]} tip_message={wanted[2]}")
    if dry_run:
        print("[dry-run] 未写文件。")
        return

    for offset, (message_id, _method, wanted) in enumerate(to_append, start=1):
        row_idx = last_data_row + offset
        ws.cell(row=row_idx, column=col["id"], value=message_id)
        for name, value in zip(LIMITER_COLUMNS[1:], wanted):
            ws.cell(row=row_idx, column=col[name], value=value)
    wb.save(LIMITER_XLSX)

    check_ws = only_sheet(openpyxl.load_workbook(LIMITER_XLSX), LIMITER_XLSX)
    after = {}
    for row_idx in range(DATA_BEGIN_ROW, check_ws.max_row + 1):
        value = check_ws.cell(row=row_idx, column=col["id"]).value
        if value is None or value == "":
            continue
        after[int(value)] = tuple(check_ws.cell(row=row_idx, column=col[name]).value for name in LIMITER_COLUMNS[1:])
    expected = dict(existing)
    expected.update({message_id: wanted for message_id, _m, wanted in to_append})
    if after != expected:
        fail(f"保存后的数据行与预期不符(应为原有 {len(existing)} 行原样 + 新增 {len(to_append)} 行)—— "
             f"立刻 git checkout -- {LIMITER_XLSX.relative_to(REPO_ROOT).as_posix()} 还原")
    print(f"[完成] 数据行 {len(existing)} -> {len(after)}。接着再跑一次导表器;"
          f"通过标准:generated/tables/messagelimiter.json 里能查到上面这些 id。"
          f"部署时按 路由服 → gate(重载 MessageLimiter)→ guild 的顺序重启(90-consistency.md Y-12)。")


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("command", choices=("new-tables", "tip-codes", "message-limiter"))
    parser.add_argument("--dry-run", action="store_true", help="只打印将要做什么,不写文件")
    args = parser.parse_args()
    {"new-tables": cmd_new_tables, "tip-codes": cmd_tip_codes, "message-limiter": cmd_message_limiter}[args.command](args.dry_run)


if __name__ == "__main__":
    # Windows 默认控制台编码(cp936 / cp1252)打不出部分字符时会在 print 上抛 UnicodeEncodeError,
    # 而那时文件可能已经写完 —— 统一成 utf-8,免得"改成功了却以非 0 退出"误导人。
    for stream in (sys.stdout, sys.stderr):
        if hasattr(stream, "reconfigure"):
            stream.reconfigure(encoding="utf-8", errors="replace")
    main()
