#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""friend 移植收尾要对两张 xlsx 做的两处改动,写成可重复执行的脚本。

为什么是脚本而不是手改:
  - data/tip/Tip.xlsx 同时有多个会话在写,它是二进制、不能 3-way 合并,丢行时**零报错**
    (导表器照常跑通,只是少发几个号)。唯一安全的写法是 openpyxl 从磁盘重新载入 →
    只动自己那一格 / 那一段 → save → 读回来按段核对(docs/design/friend-handoff-20260920.md §5.4)。
  - data/MessageLimiter.xlsx 按**消息号**登记档位,而消息号要等全量 proto-gen 之后才定下来
    (新方法落到哪个空号由 Go map 迭代顺序决定,每次生成都可能不同)。所以这里按**方法名**
    去 proto/message_id.txt 里查号,而不是把号写死在文档里让人抄。

两个子命令都**幂等**:已经是目标状态就只报告、不改文件;发现与预期不符的现状就中止,绝不覆盖。

用法(仓库根目录执行;需要 Python 3 + openpyxl,`pip install openpyxl`):

    python tools/scripts/friend_xlsx_patch.py tip-text           # 导表器之前跑
    python tools/scripts/friend_xlsx_patch.py message-limiter    # 全量 proto-gen 之后、第二次导表之前跑

加 --dry-run 只打印将要做什么,不写文件。
"""

from __future__ import annotations

import argparse
import sys
from pathlib import Path

try:
    import openpyxl
except ImportError:  # pragma: no cover - 环境问题,给一句能照着做的话
    sys.exit("缺少 openpyxl:请先执行  pip install openpyxl")

REPO_ROOT = Path(__file__).resolve().parents[2]

TIP_XLSX = REPO_ROOT / "data" / "tip" / "Tip.xlsx"
TIP_GROUP_HEADER_PREFIX = "//"
TIP_FRIEND_GROUP = "//friend_error"
TIP_CODE_NAME = "FriendBlocked"
# 旧文案点明了拉黑方向。AddFriend / AcceptFriend 的两个拉黑方向共用这一个码:
# 告诉申请人"是对方拉黑了你"等于把别人的黑名单变成可探测接口,而在"我拉黑了对方"的方向上
# 这句话本身也是错的(go/friend/internal/data/friend_repo.go 的 ErrBlocked 注释)。用户 2026-09-20 拍板改中性文案。
TIP_OLD_TEXT = "对方已将你拉黑"
TIP_NEW_TEXT = "无法添加该玩家为好友"

MESSAGE_ID_TXT = REPO_ROOT / "proto" / "message_id.txt"
LIMITER_XLSX = REPO_ROOT / "data" / "MessageLimiter.xlsx"
LIMITER_COLUMNS = ("id", "max_requests", "time_window", "tip_message")
# 数据块从第 6 行开始(导表器的 data_begin_row;2–5 行是留白的表头区)。
LIMITER_DATA_BEGIN_ROW = 6
FRIEND_SERVICE = "ClientPlayerFriend"
OLD_FRIEND_SERVICE = "FriendService"
# 档位是建议值(读 10 次/秒、写 5 次/秒、推荐 1 次/秒),没有更早的仓内约定;
# time_window / tip_message 两列照抄既有行(1 秒 / 1000 = kRateLimitExceeded)。
# NotifyFriendEvent 是 S2C 下行,客户端不发,按惯例不进表。
LIMITER_TIERS = {
    "GetFriendList": 10,
    "GetPendingRequests": 10,
    "ListBlocks": 10,
    "AddFriend": 5,
    "AcceptFriend": 5,
    "RejectFriend": 5,
    "RemoveFriend": 5,
    "Block": 5,
    "Unblock": 5,
    "RecommendFriends": 1,
}
LIMITER_TIME_WINDOW = 1
LIMITER_TIP_MESSAGE = 1000


def fail(message: str) -> "NoReturn":
    sys.exit(f"[中止] {message}")


# ── tip-text ─────────────────────────────────────────────────────────────────


def tip_group_counts(ws) -> dict[str, int]:
    """按组头统计每段的码数。改动前后各算一次,用来证明没有碰到别人的段。"""
    counts: dict[str, int] = {}
    current = None
    for (name,) in ws.iter_rows(min_row=1, max_col=1, values_only=True):
        if not isinstance(name, str) or not name.strip():
            continue
        name = name.strip()
        if name.startswith(TIP_GROUP_HEADER_PREFIX):
            current = name.split()[0]
            counts.setdefault(current, 0)
        elif current is not None:
            counts[current] += 1
    return counts


def find_tip_row(ws) -> int:
    """返回 FriendBlocked 所在行号,并确认它落在 friend_error 段内、全表只出现一次。"""
    current = None
    hits: list[tuple[int, str | None]] = []
    for row_idx, (name,) in enumerate(ws.iter_rows(min_row=1, max_col=1, values_only=True), start=1):
        if not isinstance(name, str):
            continue
        name = name.strip()
        if name.startswith(TIP_GROUP_HEADER_PREFIX):
            current = name.split()[0]
        elif name == TIP_CODE_NAME:
            hits.append((row_idx, current))
    if len(hits) != 1:
        fail(f"{TIP_XLSX.name} 里 A 列等于 {TIP_CODE_NAME!r} 的行应当恰好 1 行,实际 {len(hits)} 行:{hits}")
    row_idx, group = hits[0]
    if group != TIP_FRIEND_GROUP:
        fail(f"{TIP_CODE_NAME} 在第 {row_idx} 行,但它所在的段是 {group!r} 而不是 {TIP_FRIEND_GROUP!r}")
    return row_idx


def cmd_tip_text(dry_run: bool) -> None:
    wb = openpyxl.load_workbook(TIP_XLSX)
    if len(wb.sheetnames) != 1:
        fail(f"{TIP_XLSX.name} 预期只有一个 sheet,实际 {wb.sheetnames}")
    ws = wb[wb.sheetnames[0]]

    before_counts = tip_group_counts(ws)
    before_rows = ws.max_row
    row_idx = find_tip_row(ws)
    cell = ws.cell(row=row_idx, column=2)
    current = cell.value.strip() if isinstance(cell.value, str) else cell.value

    if current == TIP_NEW_TEXT:
        print(f"[已是目标状态] B{row_idx} = {TIP_NEW_TEXT!r},不改文件。")
        return
    if current != TIP_OLD_TEXT:
        fail(f"B{row_idx} 现值是 {current!r},既不是旧文案 {TIP_OLD_TEXT!r} 也不是新文案 —— "
             f"有人改过这一格,先确认再决定,本脚本不覆盖。")

    print(f"B{row_idx}: {TIP_OLD_TEXT!r} -> {TIP_NEW_TEXT!r}")
    if dry_run:
        print("[dry-run] 未写文件。")
        return

    cell.value = TIP_NEW_TEXT
    wb.save(TIP_XLSX)

    # 读回来核对:行数、各段码数都不许变,只有这一格变了。
    check = openpyxl.load_workbook(TIP_XLSX)
    check_ws = check[check.sheetnames[0]]
    if check_ws.max_row != before_rows:
        fail(f"保存后行数从 {before_rows} 变成 {check_ws.max_row} —— 立刻 git checkout -- {TIP_XLSX.relative_to(REPO_ROOT)} 还原")
    after_counts = tip_group_counts(check_ws)
    if after_counts != before_counts:
        fail(f"保存后各段码数变了:{before_counts} -> {after_counts} —— 立刻还原该文件")
    if check_ws.cell(row=row_idx, column=2).value != TIP_NEW_TEXT:
        fail("保存后读回的文案不是新文案 —— 立刻还原该文件")
    print(f"[完成] 共 {sum(after_counts.values())} 个码、{len(after_counts)} 个段,段计数与改动前一致;"
          f"{TIP_FRIEND_GROUP} 段 {after_counts.get(TIP_FRIEND_GROUP)} 个码。")


# ── message-limiter ──────────────────────────────────────────────────────────


def load_message_ids() -> dict[str, int]:
    ids: dict[str, int] = {}
    for line_no, line in enumerate(MESSAGE_ID_TXT.read_text(encoding="utf-8").splitlines(), start=1):
        line = line.strip()
        if not line:
            continue
        number, sep, key = line.partition("=")
        if not sep or not number.isdigit():
            fail(f"{MESSAGE_ID_TXT.name}:{line_no} 不是 `<号>=<服务名+方法名>` 形状:{line!r}")
        ids[key] = int(number)
    return ids


def resolve_friend_ids() -> dict[str, int]:
    ids = load_message_ids()
    stale = sorted(key for key in ids if key.startswith(OLD_FRIEND_SERVICE))
    if stale:
        fail(f"{MESSAGE_ID_TXT.name} 里还有旧服务名 {stale} —— 全量 proto-gen 还没跑(或没跑成功),"
             f"消息号未定,现在不能填档位。")
    resolved: dict[str, int] = {}
    missing = []
    for method in LIMITER_TIERS:
        key = FRIEND_SERVICE + method
        if key in ids:
            resolved[method] = ids[key]
        else:
            missing.append(key)
    if missing:
        fail(f"{MESSAGE_ID_TXT.name} 里找不到 {missing} —— 全量 proto-gen 还没跑,或 friend.proto 的方法名变了。")
    return resolved


def limiter_header_columns(ws) -> dict[str, int]:
    header = [c.value.strip() if isinstance(c.value, str) else c.value for c in ws[1]]
    columns = {}
    for name in LIMITER_COLUMNS:
        if name not in header:
            fail(f"{LIMITER_XLSX.name} 第 1 行找不到列 {name!r},实际表头 {header}")
        columns[name] = header.index(name) + 1
    return columns


def cmd_message_limiter(dry_run: bool) -> None:
    friend_ids = resolve_friend_ids()

    wb = openpyxl.load_workbook(LIMITER_XLSX)
    if len(wb.sheetnames) != 1:
        fail(f"{LIMITER_XLSX.name} 预期只有一个 sheet,实际 {wb.sheetnames}")
    ws = wb[wb.sheetnames[0]]
    col = limiter_header_columns(ws)

    existing: dict[int, tuple] = {}
    last_data_row = LIMITER_DATA_BEGIN_ROW - 1
    for row_idx in range(LIMITER_DATA_BEGIN_ROW, ws.max_row + 1):
        value = ws.cell(row=row_idx, column=col["id"]).value
        if value is None or value == "":
            continue
        if not isinstance(value, (int, float)) or int(value) != value:
            fail(f"{LIMITER_XLSX.name} 第 {row_idx} 行的 id 不是整数:{value!r}")
        message_id = int(value)
        if message_id in existing:
            fail(f"{LIMITER_XLSX.name} 里消息号 {message_id} 出现了两次(第 {row_idx} 行)—— 先修表")
        existing[message_id] = tuple(ws.cell(row=row_idx, column=col[name]).value for name in LIMITER_COLUMNS[1:])
        last_data_row = row_idx

    to_append = []
    for method, max_requests in LIMITER_TIERS.items():
        message_id = friend_ids[method]
        wanted = (max_requests, LIMITER_TIME_WINDOW, LIMITER_TIP_MESSAGE)
        if message_id in existing:
            if existing[message_id] != wanted:
                # 最可能的原因:这个号以前属于别的方法(被释放的旧号易主),表里留着旧方法的档位。
                fail(f"消息号 {message_id}({FRIEND_SERVICE}{method})在表里已有一行 {existing[message_id]},"
                     f"与要写的 {wanted} 不同 —— 先确认那一行是谁的,本脚本不覆盖。")
            continue
        to_append.append((message_id, method, wanted))

    if not to_append:
        print(f"[已是目标状态] {len(LIMITER_TIERS)} 个 friend C2S 消息号都已在表里,不改文件。")
        return

    print(f"现有数据行 {len(existing)} 条,末行是第 {last_data_row} 行;将追加 {len(to_append)} 行:")
    for offset, (message_id, method, wanted) in enumerate(to_append, start=1):
        print(f"  第 {last_data_row + offset} 行: id={message_id:<4} {FRIEND_SERVICE}{method:<20} "
              f"max_requests={wanted[0]} time_window={wanted[1]} tip_message={wanted[2]}")
    if dry_run:
        print("[dry-run] 未写文件。")
        return

    for offset, (message_id, _method, wanted) in enumerate(to_append, start=1):
        row_idx = last_data_row + offset
        ws.cell(row=row_idx, column=col["id"], value=message_id)
        for name, value in zip(LIMITER_COLUMNS[1:], wanted):
            ws.cell(row=row_idx, column=col[name], value=value)
    wb.save(LIMITER_XLSX)

    check = openpyxl.load_workbook(LIMITER_XLSX)
    check_ws = check[check.sheetnames[0]]
    check_ids = [
        int(v) for (v,) in check_ws.iter_rows(min_row=LIMITER_DATA_BEGIN_ROW, min_col=col["id"],
                                              max_col=col["id"], values_only=True)
        if v is not None and v != ""
    ]
    expected = len(existing) + len(to_append)
    if len(check_ids) != expected or len(set(check_ids)) != expected:
        fail(f"保存后数据行应为 {expected} 条且 id 唯一,实际 {len(check_ids)} 条 / {len(set(check_ids))} 个不同 id —— 立刻还原该文件")
    print(f"[完成] 数据行 {len(existing)} -> {expected}。接着再跑一次导表器;"
          f"通过标准:generated/tables/messagelimiter.json 里能查到上面这些 id,manifest 的 MessageLimiter rows = {expected}。")


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("command", choices=("tip-text", "message-limiter"))
    parser.add_argument("--dry-run", action="store_true", help="只打印将要做什么,不写文件")
    args = parser.parse_args()
    if args.command == "tip-text":
        cmd_tip_text(args.dry_run)
    else:
        cmd_message_limiter(args.dry_run)


if __name__ == "__main__":
    # Windows 默认控制台编码(cp936 / cp1252)打不出部分字符时会在 print 上抛 UnicodeEncodeError,
    # 而那时文件可能已经写完 —— 统一成 utf-8,免得"改成功了却以非 0 退出"误导人。
    for stream in (sys.stdout, sys.stderr):
        if hasattr(stream, "reconfigure"):
            stream.reconfigure(encoding="utf-8", errors="replace")
    main()
